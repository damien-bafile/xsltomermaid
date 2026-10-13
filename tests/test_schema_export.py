"""Tests for exporting the selected tables and columns (CSV, JSON, Excel)."""

from __future__ import annotations

import csv
import json

from xsltomermaid.excel_to_mermaid import (
    EXPECTED_HEADERS,
    build_schema,
    filter_schema,
    read_rows,
)
from xsltomermaid.make_sample import ROWS, write_sample
from xsltomermaid.schema_export import (
    DEFAULT_FIELDS,
    TableExportOptions,
    export_schema,
    relationships_path,
    sheet_name,
)


def _schema():
    return build_schema(dict(zip(EXPECTED_HEADERS, row)) for row in ROWS)


def test_csv_round_trips_through_the_reader(tmp_path):
    schema = _schema()
    path = tmp_path / "out.csv"
    written = export_schema(schema, str(path), TableExportOptions(format="csv"))
    assert written == [str(path), relationships_path(str(path))]
    with open(path, encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == sum(len(t.columns) for t in schema.tables)
    assert list(rows[0]) == [
        "SchemaName", "TableName", "ColumnOrder", "ColumnName", "DataType", "Length",
        "Precision", "Scale", "IsNullable", "IsIdentity", "IsComputed", "IsPrimaryKey",
        "ForeignKeyReference", "DefaultValue", "ComputedDefinition", "Collation",
        "Description",
    ]
    again = build_schema(rows)
    assert [t.name for t in again.tables] == [t.name for t in schema.tables]
    assert len(again.relationships) == len(schema.relationships)


def test_csv_options(tmp_path):
    path = tmp_path / "out.csv"
    opts = TableExportOptions(
        format="csv",
        fields=["table", "name", "is_primary_key"],
        spreadsheet_headers=False,
        bool_style="yes_no",
        delimiter=";",
        excel_bom=False,
        include_relationships=False,
    )
    assert export_schema(_schema(), str(path), opts) == [str(path)]
    text = path.read_text(encoding="utf-8")
    assert not text.startswith("﻿")
    lines = text.splitlines()
    assert lines[0] == "Table;Column;Primary key"
    assert all(line.split(";")[2] in ("Yes", "No") for line in lines[1:])


def test_json_nested_and_flat(tmp_path):
    schema = filter_schema(_schema(), ["Customer", "Order"])
    path = tmp_path / "out.json"
    export_schema(schema, str(path), TableExportOptions(format="json"))
    data = json.loads(path.read_text(encoding="utf-8"))
    assert [t["name"] for t in data["tables"]] == [t.name for t in schema.tables]
    first = data["tables"][0]["columns"][0]
    assert "TableName" not in first  # the table's, not each column's
    assert isinstance(first["IsNullable"], bool)
    assert isinstance(first["ColumnOrder"], int)
    assert data["relationships"]

    export_schema(schema, str(path), TableExportOptions(format="json", nested=False))
    flat = json.loads(path.read_text(encoding="utf-8"))
    assert len(flat["columns"]) == sum(len(t.columns) for t in schema.tables)
    assert flat["columns"][0]["TableName"]


def test_xlsx_single_sheet_reloads(tmp_path):
    from openpyxl import load_workbook

    schema = _schema()
    path = tmp_path / "out.xlsx"
    export_schema(schema, str(path), TableExportOptions(format="xlsx"))
    wb = load_workbook(path)
    assert wb.sheetnames == ["Columns", "Tables", "Relationships"]
    assert wb["Columns"]["A1"].value == "SchemaName"
    again = build_schema(read_rows(str(path)))
    assert [t.name for t in again.tables] == [t.name for t in schema.tables]


def test_xlsx_sheet_per_table_and_formula_text(tmp_path):
    from openpyxl import load_workbook

    schema = filter_schema(_schema(), ["Customer"])
    schema.tables[0].columns[0].default_value = "=getdate()"
    path = tmp_path / "out.xlsx"
    opts = TableExportOptions(
        format="xlsx", sheet_per_table=True, summary_sheet=False,
        fields=DEFAULT_FIELDS,
    )
    export_schema(schema, str(path), opts)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Customer"]  # no links within one table
    col = DEFAULT_FIELDS.index("default_value") + 1
    cell = wb["Customer"].cell(row=2, column=col)
    assert cell.value == "=getdate()" and cell.data_type == "s"


def test_sheet_names_are_valid_and_unique():
    taken: set[str] = set()
    assert sheet_name("a/b:c", taken) == "a_b_c"
    long = "x" * 40
    first = sheet_name(long, taken)
    second = sheet_name(long, taken)
    assert len(first) == 31 and len(second) == 31 and first != second


def test_sample_file_exports(tmp_path):
    sample = tmp_path / "sample.xlsx"
    write_sample(str(sample))
    schema = build_schema(read_rows(str(sample)))
    for kind in ("csv", "json", "xlsx"):
        out = tmp_path / f"out.{kind}"
        assert export_schema(schema, str(out), TableExportOptions(format=kind))
