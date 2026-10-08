"""Headless GUI test: build the window, load the sample, and screenshot it.

Runs Qt with the offscreen platform so it works in CI with no display. The test
asserts the window actually painted something (a PNG of a sensible size).

Run: QT_QPA_PLATFORM=offscreen uv run pytest tests/test_screenshot.py
"""

from __future__ import annotations

import os
import subprocess
import sys

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
# Keep remembered settings (export format, …) out of the user's real profile.
os.environ.setdefault(
    "XSLTOMERMAID_SETTINGS",
    os.path.join(__import__("tempfile").mkdtemp(prefix="xsltomermaid_test_"), "settings.ini"),
)
# Needed for the WebEngine (Chromium) diagram render to run headless / as root.
os.environ.setdefault("QTWEBENGINE_DISABLE_SANDBOX", "1")
os.environ.setdefault(
    "QTWEBENGINE_CHROMIUM_FLAGS", "--no-sandbox --disable-gpu --in-process-gpu"
)

import struct
import tempfile
import xml.etree.ElementTree as ET

import pytest

# Skip cleanly if the GUI stack (PySide6 + system libs) isn't importable.
pytest.importorskip("PySide6")

from PySide6.QtWidgets import QApplication  # noqa: E402

from xsltomermaid import main as app_module  # noqa: E402
from xsltomermaid.make_sample import ROWS  # noqa: E402


def test_fit_scale_caps_large_exports():
    """A huge diagram at high scale is clamped to stay within the size caps."""
    from xsltomermaid.diagram_view import MAX_PNG_DIM, MAX_PNG_PIXELS, _fit_scale

    # Small diagram: requested scale is kept as-is.
    assert _fit_scale(400, 300, 2.0) == 2.0

    # Large diagram at 4x would be ~768 MP; must clamp under both caps.
    eff = _fit_scale(8000, 6000, 4.0)
    assert eff < 4.0
    assert 8000 * eff <= MAX_PNG_DIM + 1
    assert 6000 * eff <= MAX_PNG_DIM + 1
    assert (8000 * eff) * (6000 * eff) <= MAX_PNG_PIXELS * 1.01

    # Degenerate sizes don't blow up.
    assert _fit_scale(0, 0, 3.0) == 3.0


def test_render_style_defaults_to_left_to_right():
    from xsltomermaid.diagram_view import RenderStyle

    assert RenderStyle().layout_direction == "LR"


def test_diagram_options_bar_constructs_headlessly():
    app = QApplication.instance() or QApplication([])
    bar = app_module.DiagramOptionsBar()

    assert bar.render_style().layout_direction == "LR"


def test_column_selector_marks_primary_and_foreign_keys():
    from xsltomermaid.excel_to_mermaid import Column, Table

    app = QApplication.instance() or QApplication([])
    selector = app_module.ColumnSelector()
    selector.set_tables(
        [
            Table(
                "dbo",
                "Orders",
                [
                    Column("dbo", "Orders", 1, "OrderID", "int", is_primary_key=True),
                    Column(
                        "dbo",
                        "Orders",
                        2,
                        "CustomerID",
                        "int",
                        foreign_key_reference="dbo.Customer.CustomerID",
                    ),
                ],
            )
        ]
    )

    table_item = selector._tree.topLevelItem(0)
    assert table_item.child(0).text(0) == "OrderID  [PK]"
    assert table_item.child(1).text(0) == "CustomerID  [FK]"
    assert table_item.child(0).toolTip(0) == "PK key column: OrderID"
    del app


def test_sql_query_dock_is_fixed_on_right_and_copyable():
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    window.show()
    app.processEvents()
    assert window._sql_dock.isHidden()
    assert not window._sql_dock.toggleViewAction().isChecked()
    window._sql_dock.toggleViewAction().trigger()
    assert window._sql_dock.isVisible()

    assert window.dockWidgetArea(window._sql_dock) == app_module.Qt.RightDockWidgetArea
    assert window._sql_dock.allowedAreas() == app_module.Qt.RightDockWidgetArea
    assert not window._sql_dock.isFloating()
    assert window._sql_dock.features() == app_module.QDockWidget.DockWidgetClosable
    for header in app_module.EXPECTED_HEADERS:
        assert f"[{header}]" in window._sql_query_view.toPlainText()

    window.copy_sql_query()
    assert app.clipboard().text() == app_module.SQL_SERVER_SCHEMA_QUERY
    window._diagram_view.cleanup()
    window.close()
    del app


@pytest.mark.parametrize(
    "mode,keys_first,expected",
    [
        ("order", False, ["zeta", "Id", "ParentId", "alpha"]),
        ("name", False, ["alpha", "Id", "ParentId", "zeta"]),
        ("name_desc", False, ["zeta", "ParentId", "Id", "alpha"]),
        ("type", False, ["alpha", "Id", "ParentId", "zeta"]),
        ("name", True, ["Id", "ParentId", "alpha", "zeta"]),
        ("name_desc", True, ["Id", "ParentId", "zeta", "alpha"]),
    ],
)
def test_column_sorting_preserves_selection_and_source(mode, keys_first, expected):
    from xsltomermaid.excel_to_mermaid import Column, Schema, Table, filter_columns, generate_mermaid

    app = QApplication.instance() or QApplication([])
    table = Table("dbo", "Example", [
        Column("dbo", "Example", 1, "zeta", "varchar"),
        Column("dbo", "Example", 2, "Id", "int", is_primary_key=True),
        Column("dbo", "Example", 3, "ParentId", "int",
               foreign_key_reference="dbo.Parent.Id"),
        Column("dbo", "Example", 4, "alpha", "bit"),
    ])
    selector = app_module.ColumnSelector()
    selector.set_tables([table])
    selector._tree.topLevelItem(0).child(0).setCheckState(0, app_module.Qt.Unchecked)
    selector._filter.setText("Id")
    selector._sort.setCurrentIndex(selector._sort.findData(mode))
    selector._keys_first.setChecked(keys_first)

    assert [c.name for c in selector.ordered_columns(table)] == expected
    children = list(selector._children(selector._tree.topLevelItem(0)))
    assert [c.text(0).split("  [")[0] for c in children] == expected
    assert next(c for c in children if c.text(0) == "zeta").checkState(0) == app_module.Qt.Unchecked
    assert next(c for c in children if c.text(0) == "alpha").isHidden()
    assert selector.excluded_pairs() == {("example", "zeta")}
    schema = selector.sorted_schema(filter_columns(Schema([table], []), selector.excluded_pairs()))
    assert [c.name for c in schema.tables[0].columns] == [n for n in expected if n != "zeta"]
    mermaid = generate_mermaid(schema)
    positions = [mermaid.index(f" {name}") for name in expected if name != "zeta"]
    assert positions == sorted(positions)
    assert [c.name for c in table.columns] == ["zeta", "Id", "ParentId", "alpha"]
    del app


def test_extracted_data_header_sorting_is_numeric_and_non_destructive():
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    rows = [
        [str(order) if h == "ColumnOrder" else name if h == "ColumnName" else ""
         for h in app_module.EXPECTED_HEADERS]
        for order, name in [(10, "zeta"), (2, "Alpha"), (1, "beta")]
    ]
    window._table_model.set_rows(rows)
    order_column = app_module.EXPECTED_HEADERS.index("ColumnOrder")
    name_column = app_module.EXPECTED_HEADERS.index("ColumnName")
    proxy = window._table.model()
    window._table.sortByColumn(order_column, app_module.Qt.AscendingOrder)
    assert [proxy.index(i, order_column).data() for i in range(3)] == ["1", "2", "10"]
    window._table.sortByColumn(order_column, app_module.Qt.DescendingOrder)
    assert [proxy.index(i, order_column).data() for i in range(3)] == ["10", "2", "1"]
    window._table.sortByColumn(name_column, app_module.Qt.AscendingOrder)
    assert [proxy.index(i, name_column).data() for i in range(3)] == ["Alpha", "beta", "zeta"]
    assert window._table_model._rows == rows
    window._diagram_view.cleanup()
    del app


@pytest.fixture(autouse=True)
def _isolated_settings(tmp_path, monkeypatch):
    """Each test starts from empty settings: remembered options, recent files
    and window layout must not leak from one test into the next."""
    monkeypatch.setenv("XSLTOMERMAID_SETTINGS", str(tmp_path / "settings.ini"))


def _wait_for(condition, timeout_ms: int = 3000) -> bool:
    """Pump events until ``condition()`` is true (PySide6 lacks QTest.qWaitFor)."""
    import time
    from PySide6.QtTest import QTest

    deadline = time.monotonic() + timeout_ms / 1000
    while not condition() and time.monotonic() < deadline:
        QTest.qWait(20)
    return condition()


def test_diagram_zoom_label_tracks_webengine_zoom():
    import xsltomermaid.diagram_view as diagram_view
    from PySide6.QtTest import QTest

    if not diagram_view.WEBENGINE_AVAILABLE:
        pytest.skip("PySide6 WebEngine not available")
    app = QApplication.instance() or QApplication([])
    view = diagram_view.DiagramView()
    assert view._zoom_label.text() == "100%"
    for zoom, text in [(1.25, "125%"), (0.5, "50%"), (1.0, "100%")]:
        view._view.setZoomFactor(zoom)
        # The label polls every 200 ms; under a loaded test run a fixed wait
        # can miss the tick, so wait for the condition (up to 3 s) instead.
        _wait_for(lambda: view._zoom_label.text() == text, 3000)
        assert view._zoom_label.text() == text
    view.cleanup()
    assert not view._zoom_timer.isActive()
    del app


def test_schema_to_drawio_creates_table_vertices_and_edges():
    import xsltomermaid.diagram_view as diagram_view
    from xsltomermaid.excel_to_mermaid import build_schema

    records = [dict(zip(app_module.EXPECTED_HEADERS, row)) for row in ROWS]
    schema = build_schema(records)
    drawio = diagram_view.schema_to_drawio(schema)
    root = ET.fromstring(drawio)

    assert root.tag == "mxfile"
    vertices = root.findall(".//mxCell[@vertex='1']")
    edges = root.findall(".//mxCell[@edge='1']")
    assert len(vertices) == len(schema.tables)
    assert len(edges) == len(schema.relationships)
    assert all("shape=mxgraph.er.entity" in cell.attrib.get("style", "") for cell in vertices)
    assert all("endArrow=ERmany" in cell.attrib.get("style", "") for cell in edges)
    # Every edge must carry a geometry, or Draw.io opens the saved file blank
    # (the import path is lenient, but File > Open is not).
    assert edges, "expected at least one relationship edge"
    assert all(cell.find("mxGeometry") is not None for cell in edges)


def test_svg_dimensions_handles_fractional_and_exponent_sizes():
    import xsltomermaid.diagram_view as diagram_view

    w, h = diagram_view._svg_dimensions('<svg width=".5" height="1e3"></svg>')
    assert w == 0.5
    assert h == 1000.0
    w1, h1 = diagram_view._svg_dimensions('<svg width="1e3px" height="2.5e2pt"></svg>')
    assert w1 == 1000.0
    assert h1 == 250.0
    w0, h0 = diagram_view._svg_dimensions('<svg width="0" height="0"></svg>')
    assert w0 == 0.0
    assert h0 == 0.0
    # A percentage width (Mermaid's useMaxWidth) must fall back to the viewBox,
    # not be read as 100 — otherwise the PDF fallback would size to 100px.
    wp, hp = diagram_view._svg_dimensions(
        '<svg width="100%" viewBox="0 0 3526 6531"></svg>'
    )
    assert wp == 3526.0
    assert hp == 6531.0


def test_schema_to_drawio_dark_variant():
    import xsltomermaid.diagram_view as diagram_view
    from xsltomermaid.excel_to_mermaid import Relationship, Schema, Table

    schema = Schema(
        tables=[Table("dbo", "A"), Table("dbo", "B")],
        relationships=[Relationship("A", "B", "AID")],
    )
    light = ET.fromstring(diagram_view.schema_to_drawio(schema))
    dark = ET.fromstring(diagram_view.schema_to_drawio(schema, dark=True))

    # Dark: dark page background, dark box fill + light font, light edge stroke.
    gm_d = dark.find(".//mxGraphModel")
    assert gm_d.attrib.get("background") == "#1e1e1e"
    dv = [c.attrib["style"] for c in dark.findall(".//mxCell[@vertex='1']")]
    assert dv and all("fillColor=#2b2b2b" in s and "fontColor=#e8eaed" in s for s in dv)
    de = [c.attrib["style"] for c in dark.findall(".//mxCell[@edge='1']")]
    assert de and all("strokeColor=#9aa0a6" in s for s in de)

    # Light (default) is unchanged: white fill, no page background, no fontColor.
    gm_l = light.find(".//mxGraphModel")
    assert "background" not in gm_l.attrib
    lv = [c.attrib["style"] for c in light.findall(".//mxCell[@vertex='1']")]
    assert lv and all("fillColor=#ffffff" in s and "fontColor" not in s for s in lv)


def test_schema_to_drawio_includes_table_names_and_fk_labels():
    import xsltomermaid.diagram_view as diagram_view
    from xsltomermaid.excel_to_mermaid import build_schema

    records = [dict(zip(app_module.EXPECTED_HEADERS, row)) for row in ROWS]
    schema = build_schema(records)
    drawio = diagram_view.schema_to_drawio(schema)
    root = ET.fromstring(drawio)
    vertex_values = [
        cell.attrib.get("value", "") for cell in root.findall(".//mxCell[@vertex='1']")
    ]
    edge_values = [
        cell.attrib.get("value", "") for cell in root.findall(".//mxCell[@edge='1']")
    ]
    assert any("<b>Customer</b>" in value for value in vertex_values)
    assert any("<b>Order</b>" in value for value in vertex_values)
    assert "CustomerID" in edge_values


def test_schema_to_excalidraw_scene_is_valid_and_bound():
    import json

    import xsltomermaid.diagram_view as diagram_view
    from xsltomermaid.excel_to_mermaid import build_schema

    records = [dict(zip(app_module.EXPECTED_HEADERS, row)) for row in ROWS]
    schema = build_schema(records)
    scene = json.loads(diagram_view.schema_to_excalidraw(schema))

    assert scene["type"] == "excalidraw"
    by_id = {el["id"]: el for el in scene["elements"]}
    rects = [e for e in scene["elements"] if e["type"] == "rectangle"]
    texts = [e for e in scene["elements"] if e["type"] == "text"]
    arrows = [e for e in scene["elements"] if e["type"] == "arrow"]

    assert len(rects) == len(schema.tables)
    assert len(arrows) == len(schema.relationships)

    # Table names + FK labels present in the text elements.
    joined = "\n".join(t["text"] for t in texts)
    assert "Customer" in joined and "Order" in joined and "CustomerID" in joined

    # Every arrow is glued to real rectangles at both ends.
    for arrow in arrows:
        for side in ("startBinding", "endBinding"):
            target = arrow[side]["elementId"]
            assert target in by_id and by_id[target]["type"] == "rectangle"

    # Bound text/arrows are back-referenced from their rectangles, and all
    # referenced ids exist.
    for rect in rects:
        for ref in rect["boundElements"]:
            assert ref["id"] in by_id


def test_schema_to_excalidraw_relationship_matching_is_case_insensitive():
    import json

    import xsltomermaid.diagram_view as diagram_view
    from xsltomermaid.excel_to_mermaid import Relationship, Schema, Table

    schema = Schema(
        tables=[Table("dbo", "Customer"), Table("dbo", "Order")],
        relationships=[Relationship(parent_table="CUSTOMER", child_table="order", label="CustomerID")],
    )
    scene = json.loads(diagram_view.schema_to_excalidraw(schema))
    arrows = [e for e in scene["elements"] if e["type"] == "arrow"]
    assert len(arrows) == 1


def test_schema_to_excalidraw_dark_mode():
    import json

    import xsltomermaid.diagram_view as diagram_view
    from xsltomermaid.excel_to_mermaid import Relationship, Schema, Table

    schema = Schema(
        tables=[Table("dbo", "A"), Table("dbo", "B")],
        relationships=[Relationship(parent_table="A", child_table="B", label="X")],
    )
    # Excalidraw's dark theme inverts the canvas at render time, so a dark scene
    # keeps the normal colours and only flips the theme flag.
    scene = json.loads(diagram_view.schema_to_excalidraw(schema, dark=True))
    assert scene["appState"]["theme"] == "dark"
    assert all(e["strokeColor"] == "#1e1e1e" for e in scene["elements"])
    light = json.loads(diagram_view.schema_to_excalidraw(schema))
    assert light["appState"]["theme"] == "light"
    assert all(e["strokeColor"] == "#1e1e1e" for e in light["elements"])


def test_schema_to_excalidraw_spacing_grows_with_label_length():
    import json

    import xsltomermaid.diagram_view as diagram_view
    from xsltomermaid.excel_to_mermaid import Relationship, Schema, Table

    def gap(label):
        schema = Schema(
            tables=[Table("dbo", "A"), Table("dbo", "B")],
            relationships=[Relationship(parent_table="A", child_table="B", label=label)],
        )
        els = json.loads(diagram_view.schema_to_excalidraw(schema))["elements"]
        rects = {e["id"]: e for e in els if e["type"] == "rectangle"}
        a, b = rects["rect0"], rects["rect1"]
        return b["x"] - (a["x"] + a["width"])  # horizontal gap between the boxes

    # A long relationship label widens the gap so it doesn't overlap a box.
    assert gap("A_VERY_LONG_FOREIGN_KEY_COLUMN_NAME") > gap("X")


def test_schema_to_drawio_spacing_grows_with_label_length():
    import xsltomermaid.diagram_view as diagram_view
    from xsltomermaid.excel_to_mermaid import Relationship, Schema, Table

    def gap(label):
        schema = Schema(
            tables=[Table("dbo", "A"), Table("dbo", "B")],
            relationships=[Relationship(parent_table="A", child_table="B", label=label)],
        )
        root = ET.fromstring(diagram_view.schema_to_drawio(schema))
        geoms = {}
        for cell in root.findall(".//mxCell[@vertex='1']"):
            g = cell.find("mxGeometry")
            geoms[cell.attrib["id"]] = (float(g.attrib["x"]), float(g.attrib["width"]))
        # Two tables land side by side in the same grid row (ids "2" and "3").
        (ax, aw), (bx, _bw) = geoms["2"], geoms["3"]
        return bx - (ax + aw)  # horizontal gap between the boxes

    # A long relationship label widens the inter-box gap, same as Excalidraw.
    assert gap("A_VERY_LONG_FOREIGN_KEY_COLUMN_NAME") > gap("X")


def test_schema_to_drawio_relationship_matching_is_case_insensitive():
    import xsltomermaid.diagram_view as diagram_view
    from xsltomermaid.excel_to_mermaid import Relationship, Schema, Table

    schema = Schema(
        tables=[Table("dbo", "Customer"), Table("dbo", "Order")],
        relationships=[Relationship(parent_table="CUSTOMER", child_table="order", label="CustomerID")],
    )
    drawio = diagram_view.schema_to_drawio(schema)
    root = ET.fromstring(drawio)
    vertices = {
        cell.attrib["id"]: cell.attrib.get("value", "")
        for cell in root.findall(".//mxCell[@vertex='1']")
    }
    edges = root.findall(".//mxCell[@edge='1']")
    assert len(edges) == 1
    source = edges[0].attrib.get("source")
    target = edges[0].attrib.get("target")
    assert source in vertices and "<b>Customer</b>" in vertices[source]
    assert target in vertices and "<b>Order</b>" in vertices[target]


def test_schema_to_drawio_qualified_names_avoid_ambiguous_table_matches():
    import xsltomermaid.diagram_view as diagram_view
    from xsltomermaid.excel_to_mermaid import Relationship, Schema, Table

    schema = Schema(
        tables=[
            Table("dbo", "Customer"),
            Table("sales", "Customer"),
            Table("dbo", "Order"),
        ],
        relationships=[
            Relationship(parent_table="dbo.Customer", child_table="Order", label="CustomerID"),
            Relationship(parent_table="Customer", child_table="Order", label="AmbiguousCustomer"),
        ],
    )
    drawio = diagram_view.schema_to_drawio(schema)
    root = ET.fromstring(drawio)
    edge_values = [cell.attrib.get("value", "") for cell in root.findall(".//mxCell[@edge='1']")]
    assert "CustomerID" in edge_values
    assert "AmbiguousCustomer" not in edge_values


def test_schema_to_drawio_qualified_case_collision_is_ambiguous():
    import xsltomermaid.diagram_view as diagram_view
    from xsltomermaid.excel_to_mermaid import Relationship, Schema, Table

    schema = Schema(
        tables=[Table("dbo", "Customer"), Table("DBO", "customer"), Table("dbo", "Order")],
        relationships=[
            Relationship(parent_table="dbo.customer", child_table="Order", label="ShouldSkip")
        ],
    )
    drawio = diagram_view.schema_to_drawio(schema)
    root = ET.fromstring(drawio)
    edge_values = [cell.attrib.get("value", "") for cell in root.findall(".//mxCell[@edge='1']")]
    assert "ShouldSkip" not in edge_values


def _png_size(path: str) -> tuple[int, int]:
    """Return (width, height) read from a PNG header, or (0, 0) if invalid."""
    with open(path, "rb") as handle:
        header = handle.read(24)
    if len(header) < 24 or header[:8] != b"\x89PNG\r\n\x1a\n":
        return (0, 0)
    width, height = struct.unpack(">II", header[16:24])
    return (width, height)


def _ensure_sample(path: str):
    from openpyxl import Workbook

    from xsltomermaid.excel_to_mermaid import EXPECTED_HEADERS

    wb = Workbook()
    ws = wb.active
    ws.append(EXPECTED_HEADERS)
    for row in ROWS:
        ws.append(row)
    wb.save(path)


def test_window_screenshot(tmp_path):
    # Reuse a single QApplication across tests.
    app = QApplication.instance() or QApplication([])

    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))
    window.resize(1100, 760)
    assert window._options_bar.render_style().layout_direction == "LR"
    assert not window._export_btn.isHidden()
    assert window.export_formats() == ["drawio", "excalidraw", "pdf", "png", "svg"]

    out = tmp_path / "window.png"
    window.capture(str(out))

    assert out.exists(), "screenshot file was not created"
    width, height = _png_size(str(out))
    assert width > 200 and height > 200, f"screenshot looks empty: {width}x{height}"

    # Sanity: the load actually populated the model.
    assert len(window._schema.tables) == 4
    assert window._table.model().rowCount() == len(ROWS)

    del app  # keep linters quiet; app is a singleton


def test_export_drawio_checks_schema_before_prompt(monkeypatch):
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()

    info_calls = []

    def _info(*args, **kwargs):
        info_calls.append((args, kwargs))
        return app_module.QMessageBox.Ok

    def _unexpected_dialog(*args, **kwargs):
        raise AssertionError("File dialog should not be opened when schema is empty.")

    monkeypatch.setattr(app_module.QMessageBox, "information", _info)
    monkeypatch.setattr(app_module.QFileDialog, "getSaveFileName", _unexpected_dialog)
    window.set_export_format("drawio")
    window.export_diagram()
    assert info_calls, "Expected an informational prompt for empty schema."

    window._diagram_view.cleanup()
    del app


def test_large_export_warns_but_does_not_block(monkeypatch):
    from xsltomermaid.excel_to_mermaid import Schema, Table

    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()

    small = Schema(tables=[Table("dbo", "A"), Table("dbo", "B")], relationships=[])
    big = Schema(
        tables=[Table("dbo", f"T{i}") for i in range(app_module.EXPORT_WARN_TABLES + 1)],
        relationships=[],
    )

    asked = []

    def _question(*args, **kwargs):
        asked.append(args)
        return app_module.QMessageBox.No

    monkeypatch.setattr(app_module.QMessageBox, "question", _question)

    # Small export: no warning, proceeds.
    assert window._confirm_large_export("excalidraw", small) is True
    assert not asked

    # Large export: warns; declining stops it (but it's a choice, not a cap).
    assert window._confirm_large_export("excalidraw", big) is False
    assert asked, "expected a confirmation prompt for a large export"

    # Accepting the warning proceeds (no hard limit).
    monkeypatch.setattr(
        app_module.QMessageBox, "question", lambda *a, **k: app_module.QMessageBox.Yes
    )
    assert window._confirm_large_export("drawio", big) is True

    window._diagram_view.cleanup()
    del app


def test_export_rendered_format_requires_webengine(monkeypatch, tmp_path):
    app = QApplication.instance() or QApplication([])
    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))
    window._diagram_view.available = False
    window._render_selection()

    info_calls = []

    def _info(*args, **kwargs):
        info_calls.append((args, kwargs))
        return app_module.QMessageBox.Ok

    def _unexpected_dialog(*args, **kwargs):
        raise AssertionError("File dialog should not open when WebEngine is unavailable.")

    monkeypatch.setattr(app_module.QMessageBox, "information", _info)
    monkeypatch.setattr(app_module.QFileDialog, "getSaveFileName", _unexpected_dialog)
    window.set_export_format("pdf")
    window.export_diagram()
    assert info_calls, "Expected an informational prompt when WebEngine is unavailable."

    window._diagram_view.cleanup()
    del app


def test_diagram_screenshot(tmp_path):
    """Render the actual Mermaid ER diagram to PNG and SVG (needs WebEngine)."""
    import xsltomermaid.diagram_view as diagram_view

    if not diagram_view.WEBENGINE_AVAILABLE:
        pytest.skip("PySide6 WebEngine not available")

    app = QApplication.instance() or QApplication([])

    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))

    png = tmp_path / "diagram.png"
    window.capture_diagram(str(png))
    assert png.exists()
    width, height = _png_size(str(png))
    assert width > 100 and height > 100, f"diagram PNG looks empty: {width}x{height}"

    svg = tmp_path / "diagram.svg"
    window.capture_diagram(str(svg))
    text = svg.read_text(encoding="utf-8")
    assert text.lstrip().startswith("<svg")
    # The rendered diagram should mention the tables from the sample.
    assert "OrderLine" in text and "Customer" in text

    # Relationship-label boxes must carry an inline style fill. Mermaid fills
    # them with an hsl() CSS rule QtSvg can't parse (it falls back to dark and
    # the label reads as a dark redaction bar in PNG/PDF); the inline style set
    # in current_svg() overrides that. See current_svg().
    import re

    boxes = re.findall(
        r'<rect[^>]*class="[^"]*relationshipLabelBox[^"]*"[^>]*>', text
    )
    assert boxes, "expected relationship-label boxes in the rendered SVG"
    assert all("style=" in b and "fill" in b for b in boxes), (
        "relationship-label boxes need an inline style fill so QtSvg doesn't "
        "render them as dark redaction bars"
    )

    # QtSvg ignores dominant-baseline:middle, so every centred text (table
    # cells, entity titles, edge labels) would ride above its intended y. It's
    # converted to an alphabetic baseline in current_svg(), so none should
    # remain in the export. See current_svg().
    assert "dominant-baseline: middle" not in text, (
        "centred text must drop dominant-baseline:middle so QtSvg positions it "
        "correctly"
    )
    assert re.search(r'class="[^"]*entityLabel[^"]*"', text), (
        "expected entity (cell) labels in the rendered SVG"
    )
    assert re.search(r'class="[^"]*relationshipLabel[^"]*"', text), (
        "expected relationship labels in the rendered SVG"
    )

    window._diagram_view.cleanup()
    del app


def _trace_path(window, start, end, *, replace):
    """Drive the From/To pickers + replace toggle, then trace the path."""
    window._selector._path_from.setCurrentText(start)
    window._selector._path_to.setCurrentText(end)
    window._selector._path_replace.setChecked(replace)
    window._find_shortest_path()


def test_trace_path_from_pickers_adds_to_selection(tmp_path):
    app = QApplication.instance() or QApplication([])

    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))
    # Customer is off the Order→OrderLine path, so it survives an additive trace.
    window._selector.clear_selection()
    window._selector.check_tables(["Customer"])

    _trace_path(window, "Order", "OrderLine", replace=False)

    selected = set(window._selector.selected_tables())
    assert {"Customer", "Order", "OrderLine"} <= selected

    window._diagram_view.cleanup()
    del app


def test_trace_path_replace_mode_discards_prior_selection(tmp_path):
    app = QApplication.instance() or QApplication([])

    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))
    window._selector.clear_selection()
    window._selector.check_tables(["Customer"])

    _trace_path(window, "Order", "OrderLine", replace=True)

    # Replace mode wipes the prior selection: only the path remains.
    assert set(window._selector.selected_tables()) == {"Order", "OrderLine"}

    window._diagram_view.cleanup()
    del app


def test_trace_path_undo_restores_prior_selection(tmp_path):
    app = QApplication.instance() or QApplication([])

    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))
    window._selector.clear_selection()
    window._selector.check_tables(["Customer"])

    _trace_path(window, "Order", "OrderLine", replace=False)
    assert window._selector._undo_btn.isVisibleTo(window._selector)

    window._selector.undo_last_change()
    assert set(window._selector.selected_tables()) == {"Customer"}
    assert not window._selector._undo_btn.isVisibleTo(window._selector)

    window._diagram_view.cleanup()
    del app


def test_trace_path_disconnected_keeps_selection_and_offers_no_undo(tmp_path):
    app = QApplication.instance() or QApplication([])

    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))
    window._selector.clear_selection()
    window._selector.check_tables(["Customer"])

    # Same table both ends is rejected before any mutation or snapshot.
    _trace_path(window, "Customer", "Customer", replace=True)
    assert set(window._selector.selected_tables()) == {"Customer"}
    assert not window._selector._undo_btn.isVisibleTo(window._selector)

    window._diagram_view.cleanup()
    del app


def test_large_render_warning_asks_once_per_selection(monkeypatch):
    from xsltomermaid.excel_to_mermaid import Column, Schema, Table

    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    n = app_module.RENDER_WARN_LIMIT + 5  # over the warn threshold
    tables = [
        Table("dbo", f"T{i}", [Column("dbo", f"T{i}", 1, "Id", "int", is_primary_key=True)])
        for i in range(n)
    ]
    window._schema = Schema(tables=tables, relationships=[])
    window._selector.set_tables([t.name for t in tables])
    window._selector.check_all()
    # Don't actually render (keep the test fast/headless-safe).
    monkeypatch.setattr(window._diagram_view, "set_diagram", lambda *a, **k: None)

    calls = {"n": 0}

    def _q(*a, **k):
        calls["n"] += 1
        return app_module.QMessageBox.Yes

    monkeypatch.setattr(app_module.QMessageBox, "question", _q)

    window._render_selection()
    window._render_selection()  # same big selection → must not re-ask
    assert calls["n"] == 1

    # Changing the selection (still over the limit) asks again.
    window._selector.set_selected_tables([t.name for t in tables[:-1]])
    window._render_selection()
    assert calls["n"] == 2

    window._diagram_view.cleanup()
    del app


def test_empty_load_reports_no_table_data():
    from xsltomermaid.excel_to_mermaid import Schema

    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    # A file that parsed but produced no tables (no row had TableName+ColumnName).
    window._apply_loaded("weird.xlsx", [{"Foo": "bar"}], Schema(tables=[], relationships=[]), "")
    assert "no table/column data" in window._status.text().lower()

    window._diagram_view.cleanup()
    del app


def _window_with_schema(tables, rels):
    """A MainWindow with an in-memory schema loaded (no file, no render)."""
    from xsltomermaid.excel_to_mermaid import Schema, Table

    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    window._schema = Schema(tables=[Table("", n) for n in tables], relationships=rels)
    window._selector.set_tables(tables)
    return app, window


def _diamond():
    """A→B, A→C, B→D, C→D — two equally short A→D paths."""
    from xsltomermaid.excel_to_mermaid import Relationship

    return ["A", "B", "C", "D"], [
        Relationship("A", "B", "ab"),
        Relationship("A", "C", "ac"),
        Relationship("B", "D", "bd"),
        Relationship("C", "D", "cd"),
    ]


def _abcd_chain():
    from xsltomermaid.excel_to_mermaid import Relationship

    return ["A", "B", "C", "D", "E"], [
        Relationship("A", "B", "ab"),
        Relationship("B", "C", "bc"),
        Relationship("C", "D", "cd"),
    ]


def _set_direction(window, value):
    combo = window._selector._path_direction
    combo.setCurrentIndex(combo.findData(value))


def test_trace_path_via_stop_forces_route(monkeypatch):
    app, window = _window_with_schema(*_diamond())
    monkeypatch.setattr(window, "_render_selection", lambda: None)

    window._selector._path_from.setCurrentText("A")
    window._selector._path_to.setCurrentText("D")
    window._selector._path_via.setCurrentText("C")  # force the A-C-D branch
    window._selector._path_replace.setChecked(True)
    window._find_shortest_path()

    assert set(window._selector.selected_tables()) == {"A", "C", "D"}

    window._diagram_view.cleanup()
    del app


def test_trace_path_popup_lets_user_choose_between_ties(monkeypatch):
    app, window = _window_with_schema(*_diamond())
    monkeypatch.setattr(window, "_render_selection", lambda: None)

    window._selector._path_from.setCurrentText("A")
    window._selector._path_to.setCurrentText("D")
    window._selector._path_replace.setChecked(True)

    seen = {}

    def _fake_get_item(_parent, _title, _label, items, _index, _editable):
        seen["items"] = list(items)
        return (items[0], True)  # sorted → "A → B → D"

    monkeypatch.setattr(app_module.QInputDialog, "getItem", _fake_get_item)
    window._find_shortest_path()

    assert seen["items"] == ["A → B → D", "A → C → D"]
    assert set(window._selector.selected_tables()) == {"A", "B", "D"}

    window._diagram_view.cleanup()
    del app


def test_trace_path_no_directed_path_offers_direction_retry(monkeypatch):
    app, window = _window_with_schema(*_abcd_chain())
    monkeypatch.setattr(window, "_render_selection", lambda: None)

    window._selector._path_from.setCurrentText("A")
    window._selector._path_to.setCurrentText("D")
    _set_direction(window, "forward")  # child→parent: A can't reach D forwards
    window._selector._path_replace.setChecked(True)

    asked = {"count": 0}

    def _yes(*_args, **_kwargs):
        asked["count"] += 1
        return app_module.QMessageBox.Yes

    monkeypatch.setattr(app_module.QMessageBox, "question", _yes)
    window._find_shortest_path()

    # Retrying without the direction constraint finds the whole chain.
    assert asked["count"] == 1
    assert set(window._selector.selected_tables()) == {"A", "B", "C", "D"}

    window._diagram_view.cleanup()
    del app


def test_trace_path_no_path_reports_and_keeps_selection(monkeypatch):
    app, window = _window_with_schema(*_abcd_chain())  # E is isolated
    monkeypatch.setattr(window, "_render_selection", lambda: None)
    window._selector.check_tables(["A"])

    window._selector._path_from.setCurrentText("A")
    window._selector._path_to.setCurrentText("E")
    window._selector._path_replace.setChecked(False)

    told = {"count": 0}
    monkeypatch.setattr(
        app_module.QMessageBox,
        "information",
        lambda *a, **k: told.__setitem__("count", told["count"] + 1),
    )
    window._find_shortest_path()

    assert told["count"] == 1
    assert set(window._selector.selected_tables()) == {"A"}  # untouched
    assert not window._selector._undo_btn.isVisibleTo(window._selector)  # no snapshot taken

    window._diagram_view.cleanup()
    del app


def _set_related_direction(window, value):
    combo = window._selector._related_direction
    combo.setCurrentIndex(combo.findData(value))


def _hub_schema(children=12):
    """One parent H referencing `children` child tables (a fan-out hub)."""
    from xsltomermaid.excel_to_mermaid import Relationship

    kids = [f"C{i}" for i in range(children)]
    rels = [Relationship("H", kid, "fk") for kid in kids]
    return ["H", *kids], rels


def test_add_related_direction_limits_neighbours(monkeypatch):
    app, window = _window_with_schema(*_diamond())
    monkeypatch.setattr(window, "_render_selection", lambda: None)
    window._selector.clear_selection()
    window._selector.check_tables(["D"])

    # D references B and C (forward); nothing depends on D (reverse).
    _set_related_direction(window, "forward")
    window._add_related_tables()
    assert set(window._selector.selected_tables()) == {"D", "B", "C"}

    window._selector.clear_selection()
    window._selector.check_tables(["D"])
    _set_related_direction(window, "reverse")
    window._add_related_tables()
    assert set(window._selector.selected_tables()) == {"D"}  # nothing added

    window._diagram_view.cleanup()
    del app


def test_add_related_previews_and_can_be_declined(monkeypatch):
    tables, rels = _hub_schema(children=app_module.RELATED_WARN_COUNT + 2)
    app, window = _window_with_schema(tables, rels)
    monkeypatch.setattr(window, "_render_selection", lambda: None)
    window._selector.clear_selection()
    window._selector.check_tables(["H"])
    _set_related_direction(window, "reverse")  # H's dependents = all children

    asked = {"count": 0}

    def _decline(*_args, **_kwargs):
        asked["count"] += 1
        return app_module.QMessageBox.No

    monkeypatch.setattr(app_module.QMessageBox, "question", _decline)
    window._add_related_tables()

    assert asked["count"] == 1  # a big batch prompted
    assert set(window._selector.selected_tables()) == {"H"}  # declined → unchanged
    assert not window._selector._undo_btn.isVisibleTo(window._selector)

    window._diagram_view.cleanup()
    del app


def test_add_related_accepts_and_is_undoable(monkeypatch):
    n_kids = app_module.RELATED_WARN_COUNT + 2
    tables, rels = _hub_schema(children=n_kids)
    app, window = _window_with_schema(tables, rels)
    monkeypatch.setattr(window, "_render_selection", lambda: None)
    window._selector.clear_selection()
    window._selector.check_tables(["H"])
    _set_related_direction(window, "reverse")

    monkeypatch.setattr(
        app_module.QMessageBox, "question", lambda *a, **k: app_module.QMessageBox.Yes
    )
    window._add_related_tables()

    assert len(window._selector.selected_tables()) == n_kids + 1  # H + all children
    assert window._selector._undo_btn.isVisibleTo(window._selector)
    assert window._selector._undo_btn.text() == "&Undo add"

    window._selector.undo_last_change()
    assert set(window._selector.selected_tables()) == {"H"}
    assert not window._selector._undo_btn.isVisibleTo(window._selector)

    window._diagram_view.cleanup()
    del app


def test_module_import_sets_webengine_flags_for_cli_mode():
    """`--screenshot-diagram` should preconfigure WebEngine flags at import time."""
    cmd = [
        sys.executable,
        "-c",
        (
            "import os, sys; "
            "sys.argv=['main.py','sample.xlsx','--screenshot-diagram','diagram.png']; "
            "import xsltomermaid.main; "
            "print(os.environ.get('QT_QPA_PLATFORM','')); "
            "print(os.environ.get('QTWEBENGINE_DISABLE_SANDBOX','')); "
            "print(os.environ.get('QTWEBENGINE_CHROMIUM_FLAGS','')); "
        ),
    ]
    env = {
        k: v
        for k, v in os.environ.items()
        if k != "QT_QPA_PLATFORM" and not k.startswith("QTWEBENGINE_")
    }
    proc = subprocess.run(
        cmd,
        cwd=os.path.dirname(__file__),
        env=env,
        capture_output=True,
        text=True,
        check=False,
    )
    assert proc.returncode == 0, proc.stderr
    lines = [line.strip() for line in proc.stdout.splitlines()]
    assert lines[:3] == [
        "offscreen",
        "1",
        "--no-sandbox --disable-gpu --in-process-gpu",
    ]


if __name__ == "__main__":
    from pathlib import Path

    with tempfile.TemporaryDirectory() as directory:
        test_window_screenshot(Path(directory))
        print("PASS  test_window_screenshot")
    with tempfile.TemporaryDirectory() as directory:
        test_diagram_screenshot(Path(directory))
        print("PASS  test_diagram_screenshot")


def test_help_menu_update_check_reports_both_outcomes(monkeypatch):
    from xsltomermaid.updates import Release

    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    help_menu = next(
        a.menu() for a in window.menuBar().actions() if a.text() == "&Help"
    )
    labels = [a.text() for a in help_menu.actions()]
    assert labels == ["Check for &updates…", "&About"]

    shown: list[str] = []
    opened: list[str] = []
    monkeypatch.setattr(
        app_module.QMessageBox, "information", lambda *a: shown.append(a[1])
    )
    monkeypatch.setattr(
        app_module.QMessageBox, "exec", lambda box: shown.append(box.windowTitle())
    )
    monkeypatch.setattr(app_module.webbrowser, "open", opened.append)

    current = app_module.__version__
    window._on_update_found(Release(f"v{current}", f"v{current}", "https://x/same"))
    window._on_update_found(Release("v999.0.0", "v999.0.0", "https://x/new"))
    window._on_update_failed("Couldn't reach GitHub (offline).")
    assert shown == ["No updates available", "Update available", "Couldn't check for updates"]
    assert opened == []  # nothing opens unless the user clicks the button
    assert window._status.text() == "Couldn't check for updates."
    del app


# -- auto-render, stale state, undo for bulk clears --------------------------
def test_ticking_tables_auto_renders_once_after_a_pause(monkeypatch):
    from PySide6.QtTest import QTest

    app, window = _window_with_schema(*_diamond())
    calls = []
    monkeypatch.setattr(window, "_render_selection", lambda: calls.append(1))
    items = [window._selector._list.item(i) for i in range(3)]
    for item in items:  # three quick ticks…
        item.setCheckState(app_module.Qt.Checked)
    assert window._stale and calls == []
    _wait_for(lambda: bool(calls), 2000)  # …draw once, after the pause
    QTest.qWait(app_module.AUTO_RENDER_DELAY_MS + 100)
    assert calls == [1]
    window._diagram_view.cleanup()
    del app


def test_large_selection_goes_stale_instead_of_auto_rendering(monkeypatch):
    from PySide6.QtTest import QTest

    n = app_module.RENDER_WARN_LIMIT + 1
    tables, rels = _hub_schema(children=n)
    app, window = _window_with_schema(tables, rels)
    calls = []
    monkeypatch.setattr(window, "_render_selection", lambda: calls.append(1))
    window._selector._on_select_shown()  # ticks n + 1 tables
    QTest.qWait(app_module.AUTO_RENDER_DELAY_MS + 200)
    assert calls == []
    assert window._stale
    assert "Out of date" in window._render_status._text.text()
    window._diagram_view.cleanup()
    del app


def test_export_and_copy_refuse_a_stale_diagram_the_user_declined(monkeypatch):
    app, window = _window_with_schema(*_diamond())
    window._stale = True
    monkeypatch.setattr(window, "_render_selection", lambda: None)  # "declined"
    copied = []
    monkeypatch.setattr(
        app_module.QGuiApplication.clipboard(), "setText", lambda t: copied.append(t)
    )
    window.copy_mermaid()
    assert copied == []
    assert "doesn't match your selection" in window._status.text()
    window._diagram_view.cleanup()
    del app


def test_clear_tables_is_undoable():
    app, window = _window_with_schema(*_diamond())
    window._selector.check_tables(["A", "B"])
    window._selector._on_clear()
    assert window._selector.selected_tables() == []
    assert window._selector._undo_btn.text() == "&Undo clear"
    assert window._selector._undo_btn.isVisibleTo(window._selector)
    window._selector.undo_last_change()
    assert set(window._selector.selected_tables()) == {"A", "B"}
    window._diagram_view.cleanup()
    del app


def test_column_bulk_none_is_undoable():
    from xsltomermaid.excel_to_mermaid import Column, Table

    app = QApplication.instance() or QApplication([])
    selector = app_module.ColumnSelector()
    selector.set_tables([Table("dbo", "T", [
        Column("dbo", "T", 1, "Id", "int", is_primary_key=True),
        Column("dbo", "T", 2, "Name", "varchar"),
    ])])
    changes = []
    selector.changed.connect(lambda: changes.append(1))
    selector._bulk("none")
    assert selector.excluded_pairs() == {("t", "id"), ("t", "name")}
    assert selector._undo_btn.text() == "Undo none"
    selector._undo_bulk()
    assert selector.excluded_pairs() == set()
    assert not selector._undo_btn.isVisibleTo(selector)
    assert len(changes) == 2
    del app


def test_export_format_is_remembered(tmp_path, monkeypatch):
    monkeypatch.setenv("XSLTOMERMAID_SETTINGS", str(tmp_path / "s.ini"))
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    assert window._export_kind == "png"  # default
    window.set_export_format("svg")
    assert window._export_btn.text() == "&Export SVG…"
    window._diagram_view.cleanup()
    again = app_module.MainWindow()
    assert again._export_kind == "svg"
    again._diagram_view.cleanup()
    del app


def test_preview_in_browser_works_offline(tmp_path, monkeypatch):
    app, window = _window_with_schema(*_diamond())
    window._mermaid_text = "erDiagram\n    A ||--o{ B : \"ab\"\n"
    opened = []
    monkeypatch.setattr(app_module.webbrowser, "open", opened.append)
    window.preview_browser()
    from pathlib import Path
    from urllib.parse import urlparse
    from urllib.request import url2pathname

    # url2pathname handles both file:///C:/… (Windows) and file:///tmp/… (CI).
    page = Path(url2pathname(urlparse(opened[0]).path))
    html_text = page.read_text(encoding="utf-8")
    assert '<script src="mermaid.min.js">' in html_text
    assert "cdn.jsdelivr" not in html_text
    assert (page.parent / "mermaid.min.js").stat().st_size > 100_000
    window._diagram_view.cleanup()
    del app


# -- session memory, recent files, ticked-only filter, empty state -----------
def _fresh_settings(tmp_path, monkeypatch):
    monkeypatch.setenv("XSLTOMERMAID_SETTINGS", str(tmp_path / "settings.ini"))


def test_diagram_options_survive_a_restart(tmp_path, monkeypatch):
    _fresh_settings(tmp_path, monkeypatch)
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    bar = window._options_bar
    bar._orientation.setCurrentIndex(1)  # Top → Bottom
    bar._keys_only.setChecked(True)
    bar._font.setValue(16)
    bar._more.setChecked(True)
    window.close()
    window._diagram_view.cleanup()

    again = app_module.MainWindow()
    bar = again._options_bar
    assert bar.render_style().layout_direction == "TB"
    assert bar.diagram_options().keys_only
    assert bar.render_style().font_size == 16
    assert bar._more.isChecked() and bar._more_box.isVisibleTo(bar)
    again._diagram_view.cleanup()
    del app


def test_recent_files_are_remembered_newest_first(tmp_path, monkeypatch):
    _fresh_settings(tmp_path, monkeypatch)
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _ensure_sample(str(a))
    _ensure_sample(str(b))
    window.load_file(str(a))
    window.load_file(str(b))
    window.load_file(str(a))
    assert window.recent_files() == [str(a), str(b)]
    window._rebuild_recent_menu()
    labels = [x.text() for x in window._recent_menu.actions() if x.text()]
    assert labels[:2] == ["&1  a.xlsx", "&2  b.xlsx"]

    # A vanished file is dropped from the list when picked.
    os.remove(b)
    monkeypatch.setattr(app_module.QMessageBox, "information", lambda *a, **k: None)
    window._open_recent(str(b))
    assert window.recent_files() == [str(a)]
    window._diagram_view.cleanup()
    del app


def test_ticked_only_filter_shows_just_the_selection():
    app, window = _window_with_schema(*_diamond())
    sel = window._selector
    sel.set_ready(True)
    sel.check_tables(["B", "D"])
    sel._ticked_only.setChecked(True)
    shown = [i.text() for i in sel._items() if not i.isHidden()]
    assert shown == ["B", "D"]
    sel._filter.setText("d")  # combines with the text filter
    assert [i.text() for i in sel._items() if not i.isHidden()] == ["D"]
    window._diagram_view.cleanup()
    del app


def test_extracted_data_tab_explains_the_format_until_a_file_loads(tmp_path):
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    assert window._table_stack.currentWidget() is window._table_empty
    assert "ForeignKeyReference" in window._table_empty.text()
    sample = tmp_path / "s.xlsx"
    _ensure_sample(str(sample))
    window.load_file(str(sample))
    assert window._table_stack.currentWidget() is window._table
    window._diagram_view.cleanup()
    del app


def test_sql_tab_follows_the_diagram_and_its_options(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    sample = tmp_path / "s.xlsx"
    _ensure_sample(str(sample))
    window.load_file(str(sample))
    sql = window._sql_view.toPlainText()
    assert sql.startswith("SELECT TOP (100)") and "JOIN" in sql
    assert window._tabs.tabText(window._tabs.count() - 1) == "SQL"

    window._sql_root.setCurrentText("OrderLine")
    window._sql_join.setCurrentIndex(window._sql_join.findData("LEFT"))
    window._sql_top.setValue(0)
    sql = window._sql_view.toPlainText()
    assert sql.startswith("SELECT\n") and "FROM [dbo].[OrderLine]" in sql and "LEFT JOIN" in sql

    copied = []
    monkeypatch.setattr(app_module.QGuiApplication.clipboard(), "setText", copied.append)
    window.copy_sql()
    assert copied == [sql]
    window._diagram_view.cleanup()
    del app


def _mnemonic(text: str) -> str | None:
    """The Alt+letter in a Qt label ("Ticked onl&y" → "y"); "&&" is a literal &."""
    i = 0
    while True:
        i = text.find("&", i)
        if i < 0 or i + 1 >= len(text):
            return None
        if text[i + 1] == "&":
            i += 2
            continue
        return text[i + 1].lower()


def test_alt_key_mnemonics_are_unique_on_every_tab(tmp_path):
    """Two controls on screen with the same Alt+letter make Qt cycle focus
    instead of activating either, so every visible letter must be unique."""
    from PySide6.QtWidgets import QAbstractButton, QLabel

    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    sample = tmp_path / "s.xlsx"
    _ensure_sample(str(sample))
    window.load_file(str(sample))
    window.show()
    window._options_bar._more.setChecked(True)
    window._selector._path_toggle.setChecked(True)
    window.set_details_visible(True)
    menu_letters = {_mnemonic(a.text()) for a in window.menuBar().actions()}

    for index in range(window._tabs.count() * 2):
        window.show_map(index >= window._tabs.count())  # with the diagram, then the map
        window._tabs.setCurrentIndex(index % window._tabs.count())
        app.processEvents()
        seen: dict[str, str] = {letter: "menu bar" for letter in menu_letters if letter}
        for widget in window.findChildren(QAbstractButton) + window.findChildren(QLabel):
            if not widget.isVisible():
                continue
            if isinstance(widget, QLabel) and widget.buddy() is None:
                continue  # a plain label's & is just text
            letter = _mnemonic(widget.text())
            if letter is None:
                continue
            tab = window._tabs.tabText(index % window._tabs.count())
            assert letter not in seen, (
                f"Alt+{letter.upper()} on the {tab!r} tab is used by both "
                f"{seen[letter]!r} and {widget.text()!r}"
            )
            seen[letter] = widget.text()

    for menu_action in window.menuBar().actions():
        letters = [_mnemonic(a.text()) for a in menu_action.menu().actions() if a.text()]
        letters = [x for x in letters if x]
        assert len(letters) == len(set(letters)), f"duplicate mnemonic in {menu_action.text()}"
    window._diagram_view.cleanup()
    del app


# -- export: own background, names, enablement --------------------------------
def test_export_style_relights_a_dark_view_for_paper():
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    bar = window._options_bar
    bar._theme.setCurrentIndex(bar._theme.findData("dark"))
    bar._background.setCurrentIndex(bar._background.findData("#16181d"))

    def pick(key):
        next(a for a in window._export_bg_group.actions() if a.data() == key).trigger()

    assert window.export_background() == "white"  # the default
    style = window._export_style()
    assert (style.theme, style.background) == ("default", "#ffffff")
    pick("transparent")
    assert window._export_style().background == "transparent"
    pick("match")
    assert window._export_style() == bar.render_style()
    window._diagram_view.cleanup()
    del app


def test_export_is_disabled_until_tables_are_ticked(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    sample = tmp_path / "s.xlsx"
    _ensure_sample(str(sample))
    window.load_file(str(sample))
    assert window._export_btn.isEnabled()
    window._selector._on_clear()
    window._render_selection()
    assert not window._export_btn.isEnabled()
    assert not window._preview_btn.isEnabled()
    assert "Tick some tables" in window._export_btn.toolTip()
    window._diagram_view.cleanup()
    del app


def test_save_dialog_defaults_to_the_spreadsheet_name_and_last_folder(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    window._loaded_name = "orders_schema.xlsx"
    asked = []
    out_dir = tmp_path / "exports"
    out_dir.mkdir()

    def fake_dialog(_parent, _title, default, _filter):
        asked.append(default)
        return str(out_dir / "picked.png"), ""

    monkeypatch.setattr(app_module.QFileDialog, "getSaveFileName", fake_dialog)
    window._save_path("t", ".png", "PNG (*.png)")
    window._save_path("t", ".svg", "SVG (*.svg)")
    assert asked[0] == "orders_schema.png"
    assert asked[1] == os.path.join(str(out_dir), "orders_schema.svg")
    window._diagram_view.cleanup()
    del app


def test_white_export_of_a_dark_diagram_uses_light_colours(tmp_path, monkeypatch):
    import xsltomermaid.diagram_view as diagram_view

    if not diagram_view.WEBENGINE_AVAILABLE:
        pytest.skip("PySide6 WebEngine not available")
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    bar = window._options_bar
    bar._theme.setCurrentIndex(bar._theme.findData("dark"))
    bar._background.setCurrentIndex(bar._background.findData("#16181d"))
    sample = tmp_path / "s.xlsx"
    _ensure_sample(str(sample))
    window.load_file(str(sample))
    out = tmp_path / "out.svg"
    monkeypatch.setattr(
        app_module.QFileDialog, "getSaveFileName", lambda *a, **k: (str(out), "")
    )
    window.set_export_format("svg")
    window.export_diagram()
    svg = out.read_text(encoding="utf-8").lower()
    assert "#ececff" in svg or "rgb(236, 236, 255)" in svg  # default theme's box fill
    assert "Show in folder" in window._status.text()
    assert bar.render_style().theme == "dark"  # the view itself is untouched
    window._diagram_view.cleanup()
    del app



def test_details_panel_starts_closed_and_opens_on_a_view():
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    window.show()
    assert not window._tabs.isVisible()  # closed by default: the canvas gets the width
    assert window._diagram_tab.isVisible()
    window.show_details(window._sql_tab)
    assert window._tabs.isVisible() and window._tabs.currentWidget() is window._sql_tab
    assert window._details_btn.isChecked() and window._details_action.isChecked()
    window._details_action.trigger()  # Ctrl+I closes it again
    assert not window._tabs.isVisible() and not window._details_btn.isChecked()
    window._diagram_view.cleanup()
    del app


def test_details_panel_state_is_remembered():
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    window.show_details(window._columns)
    window.close()
    window._diagram_view.cleanup()
    again = app_module.MainWindow()
    again.show()
    assert again._tabs.isVisible() and again._tabs.currentWidget() is again._columns
    again._diagram_view.cleanup()
    del app


# -- diagram ↔ list selection --------------------------------------------------
def test_entity_group_ids_parse_back_to_table_ids():
    from xsltomermaid.diagram_view import entity_id_from_group

    assert entity_id_from_group("entity-Customer-7db3a251-ee0f-5d85-b53f-42db636cc6f9") == "Customer"
    assert entity_id_from_group(
        "entity-msdyn_resourcerequirement-3ec22774-e93b-529c-bc12-1dfdb28f8e64"
    ) == "msdyn_resourcerequirement"
    assert entity_id_from_group("") is None
    assert entity_id_from_group("node-1") is None


def test_clicking_a_table_in_the_diagram_selects_it(tmp_path):
    import xsltomermaid.diagram_view as diagram_view

    if not diagram_view.WEBENGINE_AVAILABLE:
        pytest.skip("PySide6 WebEngine not available")
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    window.show()
    sample = tmp_path / "s.xlsx"
    _ensure_sample(str(sample))
    window.load_file(str(sample))
    view = window._diagram_view
    assert _wait_for(lambda: view._shell_loaded, 20000)
    view.current_svg()  # wait for Mermaid to finish drawing

    def click(kind):
        view._view.page().runJavaScript(
            "var g=document.querySelector('#container g[id^=\"entity-OrderLine-\"]');"
            f"g.dispatchEvent(new MouseEvent('{kind}', {{bubbles: true}}));"
        )

    click("click")
    assert _wait_for(lambda: window._selector._list.currentItem() is not None
                     and window._selector._list.currentItem().text() == "OrderLine", 5000)
    assert window._columns._scope_tables()[0].name == "OrderLine"
    assert not window._tabs.isVisible()  # single click doesn't open the panel
    click("dblclick")
    assert _wait_for(lambda: window._tabs.isVisible(), 5000)
    assert window._tabs.currentWidget() is window._inspector
    assert window._inspector.current_table() == "OrderLine"
    window._diagram_view.cleanup()
    del app


def test_highlighting_a_list_row_highlights_the_drawn_table(monkeypatch):
    app, window = _window_with_schema(*_diamond())
    window._entity_to_table = {"A": "A", "B": "B"}
    seen = []
    monkeypatch.setattr(window._diagram_view, "highlight_entity", seen.append)
    window._selector.set_ready(True)
    window._selector._list.setCurrentRow(1)  # "B"
    window._selector._list.setCurrentRow(3)  # "D", not drawn
    assert seen == ["B", ""]
    window._diagram_view.cleanup()
    del app



# -- the Table view (per-table inspector) ---------------------------------------
def _audit_rows():
    from tests_rows import row  # noqa: F401  (placeholder, replaced below)


def _inspector_window(tmp_path):
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    sample = tmp_path / "s.xlsx"
    _ensure_sample(str(sample))
    window.load_file(str(sample))
    return app, window


def test_table_view_shows_columns_and_both_directions_of_links(tmp_path):
    app, window = _inspector_window(tmp_path)
    window._selector.clear_selection()
    window._selector.check_tables(["Order"])
    window._render_selection()
    window._focus_table("Order")
    ins = window._inspector
    assert ins.current_table() == "Order"
    assert ins._columns.count() == 4
    groups = [ins._links.topLevelItem(i) for i in range(ins._links.topLevelItemCount())]
    assert [g.text(0) for g in groups] == ["References (1)", "Referenced by (1)"]
    out_row, in_row = groups[0].child(0), groups[1].child(0)
    assert (out_row.text(0), out_row.text(1)) == ("Customer  ·  via CustomerID", "Add")
    assert out_row.data(0, app_module.Qt.UserRole) == "Customer"
    assert (in_row.text(0), in_row.text(1)) == ("OrderLine  ·  via OrderID", "Add")
    window._diagram_view.cleanup()
    del app


def test_table_view_counts_match_the_diagram_under_keys_only(tmp_path):
    app, window = _inspector_window(tmp_path)
    window._selector.clear_selection()
    window._selector.check_tables(["Customer"])
    window._render_selection()
    window._focus_table("Customer")
    ins = window._inspector
    total = ins._columns.count()
    assert ins._meta.text().startswith(f"{total} of {total} columns drawn")
    window._options_bar._keys_only.setChecked(True)  # re-renders
    assert ins._meta.text().startswith(f"1 of {total} columns drawn (Keys only)")
    hidden = [
        ins._columns.item(i) for i in range(total)
        if ins._columns.item(i).text().endswith("hidden by Keys only")
    ]
    assert len(hidden) == total - 1
    assert all(i.checkState() == app_module.Qt.Checked for i in hidden)  # still ticked
    assert "only key columns are drawn" in window._columns._count.text()
    assert "hidden by Keys only" in window._mermaid_text
    window._diagram_view.cleanup()
    del app


def test_table_view_add_ticks_the_table_and_can_be_undone(tmp_path, monkeypatch):
    app, window = _inspector_window(tmp_path)
    monkeypatch.setattr(window, "_render_selection", lambda: None)
    window._selector.clear_selection()
    window._selector.check_tables(["Order"])
    window._focus_table("Order")
    window._inspector.add_requested.emit("Customer")
    assert set(window._selector.selected_tables()) == {"Order", "Customer"}
    assert window._selector._undo_btn.text() == "&Undo add"
    window._selector.undo_last_change()
    assert window._selector.selected_tables() == ["Order"]
    window._diagram_view.cleanup()
    del app


def test_table_view_column_ticks_change_the_diagram_columns(tmp_path):
    app, window = _inspector_window(tmp_path)
    window._focus_table("Customer")
    item = next(
        window._inspector._columns.item(i)
        for i in range(window._inspector._columns.count())
        if window._inspector._columns.item(i).data(app_module.Qt.UserRole) == "Email"
    )
    item.setCheckState(app_module.Qt.Unchecked)
    assert ("customer", "email") in window._columns.excluded_pairs()
    window._diagram_view.cleanup()
    del app


def test_table_view_folds_audit_and_ownership_links():
    from xsltomermaid.excel_to_mermaid import Column, Relationship, Table

    app, window = _window_with_schema(["systemuser", "a", "b"], [
        Relationship("systemuser", "a", "createdby", ("createdby",), ("systemuserid",)),
        Relationship("systemuser", "b", "modifiedby", ("modifiedby",), ("systemuserid",)),
        Relationship("systemuser", "b", "approver", ("approver",), ("systemuserid",)),
    ])
    window._schema.tables[0].columns.append(Column("", "systemuser", 1, "systemuserid", "guid"))
    window._focus_table("systemuser")
    group = window._inspector._links.topLevelItem(1)  # Referenced by
    assert group.text(0) == "Referenced by (3)"
    assert group.child(0).text(0) == "b  ·  via approver"
    folded = group.child(1)
    assert folded.text(0) == "Audit and ownership links (2)" and not folded.isExpanded()
    window._diagram_view.cleanup()
    del app


# -- schema map integration -------------------------------------------------------
def _big_window(n=70):
    from xsltomermaid.excel_to_mermaid import Relationship

    tables, rels = _hub_schema(children=n)
    app, window = _window_with_schema(tables, rels)
    window._map_view.set_schema(window._schema)
    return app, window


def test_map_selection_draws_as_the_diagram_and_is_undoable(monkeypatch):
    app, window = _big_window()
    window.show()
    window.show_map(True)
    assert window.map_visible() and window._map_view.schema_map() is not None
    window._selector.check_tables(["C1"])
    window._map_view.select(["H", "C3", "C4"])
    assert window._map_view.selected() == ["C3", "C4", "H"]
    rendered = []
    monkeypatch.setattr(window, "_render_selection", lambda: rendered.append(1))
    window._map_view._draw_btn.click()
    assert set(window._selector.selected_tables()) == {"H", "C3", "C4"}
    assert rendered and not window.map_visible()
    window._selector.undo_last_change()
    assert window._selector.selected_tables() == ["C1"]
    window._diagram_view.cleanup()
    del app


def test_map_words_say_tick_and_show_a_selection_the_filter_hides():
    app, window = _big_window()
    window.show()
    window.show_map(True)
    btn = window._map_view._draw_btn
    window._map_view.select(["C3", "C4"])
    window._map_view._on_selection_changed()
    assert btn.text() == "Draw these 2"
    window._selector.check_tables(["C1"])
    window._map_view.set_ticked(window._selector.selected_tables())
    assert btn.text() == "Draw these 2 (replaces 1 ticked)"
    note = window._selector._mapsel_note
    assert note.isHidden()
    window._selector._filter.setText("C1")
    assert not note.isHidden() and "2 selected on the map, all hidden" in note.text()
    note.linkActivated.emit("show")
    assert window._selector._filter.text() == "" and note.isHidden()
    assert window._selector._select_shown_btn.text() == "Tick &shown"
    assert window._selector._render_btn.text() == "D&raw ticked"
    window._diagram_view.cleanup()
    del app


def test_map_single_selection_focuses_the_table_and_filter_dims_others():
    app, window = _big_window()
    window.show()
    window.show_map(True)
    window._map_view.select(["C7"])
    window._map_view._on_selection_changed()  # as a user click would
    assert window._inspector.current_table() == "C7"
    window._selector._filter.setText("C1")
    items = window._map_view._items
    assert items["C10"].opacity() == 1.0 and items["C8"].opacity() < 0.5
    assert items["C7"].opacity() == 1.0  # selected tables are never dimmed
    window._diagram_view.cleanup()
    del app


def test_big_schema_opens_on_the_map(tmp_path):
    from xsltomermaid.excel_to_mermaid import Schema

    app, window = _big_window()
    window._apply_loaded(str(tmp_path / "big.xlsx"), [], window._schema, "")
    assert window.map_visible()
    window._diagram_view.cleanup()
    del app


def test_medium_schema_opens_on_the_map_instead_of_an_empty_canvas(tmp_path):
    app, window = _big_window(n=40)  # over the auto-draw limit, under the warning one
    window._apply_loaded(str(tmp_path / "mid.xlsx"), [], window._schema, "")
    assert window.map_visible()
    assert "Pick a cluster on the map" in window._status.text()
    window._diagram_view.cleanup()
    del app


def test_the_map_saves_as_png_and_svg(tmp_path):
    from PySide6.QtGui import QImage

    app, window = _big_window()
    window.show()
    window.show_map(True)
    png, svg = tmp_path / "map.png", tmp_path / "map.svg"
    window._map_view.export_image(str(png))
    window._map_view.export_image(str(svg))
    image = QImage(str(png))
    assert max(image.width(), image.height()) == window._map_view.EXPORT_PNG_SIDE
    assert svg.read_text(encoding="utf-8").lstrip().startswith("<?xml")
    window._diagram_view.cleanup()
    del app


def test_a_big_draw_centres_on_the_busiest_table():
    from xsltomermaid.excel_to_mermaid import filter_schema

    app, window = _big_window(n=10)
    final = filter_schema(window._schema, [t.name for t in window._schema.tables])
    window._entity_to_table = app_module.mermaid_entity_ids(final)
    assert window._entity_to_table[window._hub_entity(final)] == "H"
    window._diagram_view.cleanup()
    del app


def test_fit_stays_readable_and_small_text_is_flagged():
    from xsltomermaid import diagram_view

    assert "Math.max(force ? 0.05 : 0.75, scale)" in diagram_view._shell_html()
    app = QApplication.instance() or QApplication([])
    view = diagram_view.DiagramView()
    if view._view is None:
        return  # no WebEngine
    view._shell_loaded = True
    view._fit_scale = 0.5  # the Fit button on a big diagram: 6px text
    view._sync_small_hint()
    assert view.effective_text_px() < view.SMALL_TEXT_PX and not view._small_hint.isHidden()
    view._fit_scale = 0.75
    view._sync_small_hint()
    assert view._small_hint.isHidden()
    view.cleanup()
    del app


def test_map_export_sizes_labels_to_the_image_then_restores_them(tmp_path):
    from PySide6.QtWidgets import QGraphicsItem

    app, window = _big_window()
    window.show()
    window.show_map(True)
    labels = [label for label, _comm in window._map_view._labels]
    assert labels
    seen = {}
    real = window._map_view._write_image

    def spy(path, rect, background):
        seen["scale"] = labels[0].scale()
        seen["ignores"] = bool(labels[0].flags() & QGraphicsItem.ItemIgnoresTransformations)
        real(path, rect, background)

    window._map_view._write_image = spy
    window._map_view.export_image(str(tmp_path / "m.png"))
    assert seen["scale"] != 1.0 and not seen["ignores"]  # sized to the image
    assert labels[0].scale() == 1.0
    assert labels[0].flags() & QGraphicsItem.ItemIgnoresTransformations  # back on screen
    window._diagram_view.cleanup()
    del app


# -- smarter table list ---------------------------------------------------------
def _list_window():
    from xsltomermaid.excel_to_mermaid import Relationship

    tables = ["account", "msdyn_project", "msdyn_task", "hsl_booking", "lonely"]
    rels = [
        Relationship("msdyn_project", "msdyn_task", "project", ("project",), ("id",)),
        Relationship("account", "msdyn_project", "customer", ("customer",), ("id",)),
        Relationship("msdyn_project", "hsl_booking", "project", ("project",), ("id",)),
        Relationship("account", "lonely", "createdby", ("createdby",), ("id",)),  # audit
    ]
    app, window = _window_with_schema(tables, rels)
    window._selector.set_ready(True)
    from xsltomermaid.schema_map import link_counts
    window._selector.set_link_counts(link_counts(window._schema))
    return app, window


def _shown(sel):
    return [i.text() for i in sel._items() if not i.isHidden()]


def test_table_list_prefix_filter_and_hide_unconnected():
    app, window = _list_window()
    sel = window._selector
    prefixes = [sel._prefix.itemText(i) for i in range(sel._prefix.count())]
    assert prefixes == ["All prefixes (5)", "No prefix (2)", "msdyn_ (2)", "hsl_ (1)"]
    sel._prefix.setCurrentIndex(sel._prefix.findData("msdyn"))
    assert _shown(sel) == ["msdyn_project", "msdyn_task"]
    sel._prefix.setCurrentIndex(0)
    sel._hide_unconnected.setChecked(True)  # "lonely" has only an audit link
    assert "lonely" not in _shown(sel) and len(_shown(sel)) == 4
    window._diagram_view.cleanup()
    del app


def test_table_list_sorts_by_links_and_keeps_ticks():
    app, window = _list_window()
    sel = window._selector
    sel.check_tables(["hsl_booking"])
    sel._sort.setCurrentIndex(sel._sort.findData("links"))
    assert [i.text() for i in sel._items()][:2] == ["msdyn_project", "account"]
    assert sel.selected_tables() == ["hsl_booking"]
    item = next(i for i in sel._items() if i.text() == "msdyn_project")
    assert item.data(app_module._ROLE_LINKS) == (1, 2)
    sel.set_drawn(["msdyn_project"])
    assert item.data(app_module._ROLE_DRAWN) is True
    window._diagram_view.cleanup()
    del app


def test_setting_list_data_does_not_fire_selection_changes():
    """Link counts and drawn markers are data, not ticks: with ~1,800 rows,
    reacting to each one hung the load of a real Dynamics export."""
    app, window = _list_window()
    fired = []
    window._selector.selection_changed.connect(lambda: fired.append(1))
    from xsltomermaid.schema_map import link_counts
    window._selector.set_link_counts(link_counts(window._schema))
    window._selector.set_drawn(["account"])
    assert fired == []
    window._diagram_view.cleanup()
    del app


# -- finishing touches: hide audit links, arrow keys, sort by cluster -----------
def test_hide_audit_links_option_drops_them_from_diagram_and_sql(monkeypatch):
    from xsltomermaid.excel_to_mermaid import Column, Relationship, Table

    app, window = _window_with_schema(["systemuser", "task"], [
        Relationship("systemuser", "task", "createdby", ("createdby",), ("systemuserid",)),
        Relationship("systemuser", "task", "assignee", ("assignee",), ("systemuserid",)),
    ])
    for t, cols in (("systemuser", ["systemuserid"]), ("task", ["taskid", "createdby", "assignee"])):
        table = next(x for x in window._schema.tables if x.name == t)
        table.columns = [Column("", t, i, c, "guid") for i, c in enumerate(cols, 1)]
    window._selector.set_ready(True)
    window._selector.check_tables(["systemuser", "task"])
    window._render_selection()
    assert {r.label for r in window._drawio_schema.relationships} == {"createdby", "assignee"}
    window._options_bar._hide_audit.setChecked(True)  # re-renders via `changed`
    assert [r.label for r in window._drawio_schema.relationships] == ["assignee"]
    assert '"createdby"' not in window._mermaid_text
    assert "createdby]" not in window._sql_view.toPlainText().split(";")[0].split("FROM")[1]
    window._diagram_view.cleanup()
    del app


def test_arrow_keys_move_between_tables_and_enter_opens_one(tmp_path):
    import xsltomermaid.diagram_view as diagram_view

    if not diagram_view.WEBENGINE_AVAILABLE:
        pytest.skip("PySide6 WebEngine not available")
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    window.show()
    sample = tmp_path / "s.xlsx"
    _ensure_sample(str(sample))
    window.load_file(str(sample))
    view = window._diagram_view
    assert _wait_for(lambda: view._shell_loaded, 20000)
    view.current_svg()

    def key(name):
        view._view.page().runJavaScript(
            f"document.dispatchEvent(new KeyboardEvent('keydown', {{key: '{name}'}}));"
        )

    current = lambda: (window._selector._list.currentItem().text()  # noqa: E731
                       if window._selector._list.currentItem() else "")
    key("ArrowRight")  # nothing selected: the first table
    assert _wait_for(lambda: current() != "", 5000)
    first = current()
    key("ArrowRight")  # Left → Right layout: the next table along
    assert _wait_for(lambda: current() not in ("", first), 5000)
    key("Enter")
    assert _wait_for(lambda: window._tabs.isVisible(), 5000)
    assert window._tabs.currentWidget() is window._inspector
    assert window._inspector.current_table() == current()
    window._diagram_view.cleanup()
    del app


def test_table_list_sorts_by_cluster():
    app, window = _list_window()
    sel = window._selector
    sel.cluster_provider = lambda: {
        "msdyn_project": (0, "msdyn_project"), "msdyn_task": (0, "msdyn_project"),
        "account": (1, "account"), "hsl_booking": (0, "msdyn_project"),
    }
    sel._sort.setCurrentIndex(sel._sort.findData("cluster"))
    order = [i.text() for i in sel._items()]
    assert order == ["msdyn_project", "hsl_booking", "msdyn_task", "account", "lonely"]
    assert "cluster: msdyn_project" in next(i for i in sel._items() if i.text() == "msdyn_task").toolTip()
    window._diagram_view.cleanup()
    del app


def test_headless_capture_of_a_big_schema_includes_the_map(tmp_path):
    app, window = _big_window()
    window._apply_loaded(str(tmp_path / "big.xlsx"), [], window._schema, "")
    assert window.map_visible() and window._map_view.schema_map() is None  # never shown
    window.capture(str(tmp_path / "w.png"))
    assert window._map_view.schema_map() is not None
    assert len(window._map_view._items) > 50
    window._diagram_view.cleanup()
    del app


# -- critique run 3 fixes ----------------------------------------------------------
def _audit_heavy_window(tmp_path):
    """A schema where audit links dominate, like a Dynamics export."""
    from xsltomermaid.excel_to_mermaid import Column, Relationship, Schema, Table

    names = ["systemuser"] + [f"t{i}" for i in range(70)]
    tables = [Table("dbo", n, [Column("dbo", n, 1, f"{n}id", "guid", is_primary_key=True)]
                    + [Column("dbo", n, j + 2, f"c{j}", "int") for j in range(60)]
                    + [Column("dbo", n, 70, "createdby", "guid", foreign_key_reference="dbo.systemuser.systemuserid")])
              for n in names]
    rels = [Relationship("systemuser", f"t{i}", "createdby", ("createdby",), ("systemuserid",)) for i in range(70)]
    rels += [Relationship(f"t{i}", f"t{i+1}", "parent", ("parent",), (f"t{i}id",)) for i in range(0, 69, 2)]
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    window._apply_loaded(str(tmp_path / "dyn.xlsx"), [], Schema(tables, rels), "")
    return app, window


def test_audit_links_are_hidden_automatically_when_they_dominate(tmp_path):
    app, window = _audit_heavy_window(tmp_path)
    assert window._audit_share > 0.3
    assert window._options_bar.hide_audit_links()
    assert window._map_view._hide.isChecked()  # one setting, both checkboxes
    window._options_bar._hide_audit.setChecked(False)
    assert not window._map_view._hide.isChecked()
    window._diagram_view.cleanup()
    del app


def test_drawing_wide_tables_turns_keys_only_on_with_an_undo(tmp_path):
    app, window = _audit_heavy_window(tmp_path)
    window._draw_from_map(["t0", "t1", "t2"])
    assert window._options_bar._keys_only.isChecked()
    assert "Show all columns" in window._status.text()
    assert "createdby" not in window._mermaid_text  # audit FK column dropped too
    assert window._selector._undo_btn.text() == "&Undo draw"
    window._on_status_link("undo-keys")
    assert not window._options_bar._keys_only.isChecked()
    window._diagram_view.cleanup()
    del app


def test_state_chips_and_undo_draw_restores_keys_only(tmp_path):
    app, window = _audit_heavy_window(tmp_path)
    bar = window._options_bar
    assert not bar._audit_chip.isHidden()
    assert bar._audit_chip.text().startswith("Audit links hidden (")
    assert bar._keys_chip.isHidden()
    window._draw_from_map(["t0", "t1", "t2"])
    assert bar.keys_only_auto() and not bar._keys_chip.isHidden()
    assert "audit and system links hidden" not in window._status.text()
    window._selector.undo_last_change()  # Undo draw
    assert not bar._keys_only.isChecked() and bar._keys_chip.isHidden()
    window._draw_from_map(["t0", "t1", "t2"])
    bar._keys_chip.click()
    assert not bar._keys_only.isChecked()
    bar._keys_only.setChecked(True)  # the user's own choice: no chip, kept on undo
    assert bar._keys_chip.isHidden()
    window._selector.undo_last_change()
    assert bar._keys_only.isChecked()
    bar._audit_chip.click()
    assert not bar.hide_audit_links() and bar._audit_chip.isHidden()
    assert not window._map_view._hide.isChecked()
    window._diagram_view.cleanup()
    del app


def test_polish_contrast_names_and_zoom_shortcut(tmp_path):
    from PySide6.QtGui import QKeySequence, QPalette

    app = QApplication.instance() or QApplication([])
    app_module.apply_system_palette(app)
    assert app.palette().color(QPalette.Highlight).name() == app_module._ACCENT_FILL
    window = app_module.MainWindow()
    assert window._table.accessibleName() == "Extracted data"
    window.set_export_format("png")
    assert window._export_btn.text() == "&Export PNG…"
    assert window._export_btn.accessibleName() == "Export PNG…"
    zoom_in = next(a for a in window.menuBar().actions()[1].menu().actions()
                   if a.text() == "Zoom &in")
    assert QKeySequence("Ctrl+=") in zoom_in.shortcuts()
    border = app_module._control_border_hex(window)
    assert border in window._map_btn.styleSheet()
    window._diagram_view.cleanup()
    del app


def test_view_switch_and_wording(tmp_path):
    app, window = _audit_heavy_window(tmp_path)
    window.show()
    assert window.map_visible() and window._map_btn.isChecked()
    assert not window._options_bar._orientation.isVisible()  # diagram-only, hidden on the map
    window._diagram_view_btn.click()
    assert not window.map_visible() and window._options_bar._orientation.isVisible()
    assert window._selector._title.text() == "Tables (71)"
    assert window._selector._count.text().endswith("ticked for the diagram")
    window._map_view.select(["t3", "t4"])
    window._map_view.selection_changed.emit(["t3", "t4"])
    marked = [i.text() for i in window._selector._items() if i.data(app_module._ROLE_MAPSEL)]
    assert marked == ["t3", "t4"]
    window._diagram_view.cleanup()
    del app


def test_map_cluster_click_keyboard_and_light_colours(tmp_path):
    from PySide6.QtCore import QEvent
    from PySide6.QtGui import QKeyEvent

    app, window = _audit_heavy_window(tmp_path)
    window.show()
    mv = window._map_view
    m = mv.schema_map()
    biggest = m.communities[0]
    mv.select_cluster(0)
    assert mv.selected() == sorted(biggest.members)
    mv._view.setFocus()
    mv._on_key(QKeyEvent(QEvent.KeyPress, app_module.Qt.Key_Right, app_module.Qt.NoModifier))
    assert mv._current in m.nodes
    opened = []
    mv.table_activated.connect(lambda n, d: opened.append((n, d)))
    mv._on_key(QKeyEvent(QEvent.KeyPress, app_module.Qt.Key_Return, app_module.Qt.NoModifier))
    assert opened == [(mv._current, True)]
    colour = mv._unconnected.styleSheet()
    assert "#000000" not in colour  # the muted label kept its tone
    window._diagram_view.cleanup()
    del app


def test_cluster_sort_adds_a_header_per_cluster():
    app, window = _list_window()
    sel = window._selector
    sel.cluster_provider = lambda: {
        "msdyn_project": (0, "msdyn_project"), "msdyn_task": (0, "msdyn_project"),
        "account": (1, "account"), "hsl_booking": (0, "msdyn_project"),
    }
    sel._sort.setCurrentIndex(sel._sort.findData("cluster"))
    headers = [(i.text(), i.data(app_module._ROLE_GROUP)) for i in sel._items() if i.data(app_module._ROLE_GROUP)]
    assert headers == [("msdyn_project", "msdyn_project · 3 tables"),
                       ("account", "account · 1 tables"), ("lonely", "Unconnected · 1 tables")]
    sel._sort.setCurrentIndex(sel._sort.findData("name"))
    assert not any(i.data(app_module._ROLE_GROUP) for i in sel._items())
    window._diagram_view.cleanup()
    del app


def test_zoom_readout_includes_the_fit_scale():
    import xsltomermaid.diagram_view as diagram_view

    if not diagram_view.WEBENGINE_AVAILABLE:
        pytest.skip("PySide6 WebEngine not available")
    app = QApplication.instance() or QApplication([])
    view = diagram_view.DiagramView()
    view._on_fit_scale(0.55)
    assert view._zoom_label.text() == "55%"
    view.cleanup()
    del app


def test_saving_a_table_list_reports_write_errors(tmp_path, monkeypatch):
    app, window = _list_window()
    window._loaded_name = "x.xlsx"
    monkeypatch.setattr(app_module.QFileDialog, "getSaveFileName",
                        lambda *a, **k: (str(tmp_path / "missing" / "list.toml"), ""))
    warned = []
    monkeypatch.setattr(app_module.QMessageBox, "warning", lambda *a, **k: warned.append(a[1]))
    window.save_table_selection_toml()
    assert warned == ["Couldn't save the table list"]
    window._diagram_view.cleanup()
    del app


def test_cluster_header_rows_are_double_height_and_sorting_keeps_the_open_table():
    app, window = _list_window()
    sel = window._selector
    sel.cluster_provider = lambda: {"msdyn_project": (0, "msdyn_project"), "msdyn_task": (0, "msdyn_project")}
    window._focus_table("msdyn_task")
    assert window._inspector.current_table() == "msdyn_task"
    sel._sort.setCurrentIndex(sel._sort.findData("cluster"))
    assert window._inspector.current_table() == "msdyn_task"  # not emptied by the re-sort
    rows = list(sel._items())
    header_row = sel._list.row(next(i for i in rows if i.data(app_module._ROLE_GROUP)))
    plain_row = sel._list.row(next(i for i in rows if not i.data(app_module._ROLE_GROUP)))
    assert sel._list.sizeHintForRow(header_row) >= 2 * sel._list.sizeHintForRow(plain_row) - 1
    window._diagram_view.cleanup()
    del app
