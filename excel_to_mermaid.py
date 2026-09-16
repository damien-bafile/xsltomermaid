"""Core logic: turn a database-schema spreadsheet into a Mermaid ER diagram.

The spreadsheet is expected to have one row per *column* of a database, with the
headers shown in the reference screenshot:

    SchemaName, TableName, ColumnOrder, ColumnName, DataType, Length, Precision,
    Scale, IsNullable, IsIdentity, IsComputed, IsPrimaryKey, ForeignKeyReference,
    DefaultValue, ComputedDefinition, Collation, Description

This module is deliberately free of any GUI or heavy dependencies for the parsing
and generation steps so it can be unit-tested on its own. ``openpyxl`` is only
imported inside :func:`read_rows` so the rest of the module works without it.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Callable, Iterable

# A progress callback receives a fraction in the range 0.0 .. 1.0.
ProgressCallback = Callable[[float], None]

# The canonical headers, in the order they appear in the reference sheet.
EXPECTED_HEADERS: list[str] = [
    "SchemaName",
    "TableName",
    "ColumnOrder",
    "ColumnName",
    "DataType",
    "Length",
    "Precision",
    "Scale",
    "IsNullable",
    "IsIdentity",
    "IsComputed",
    "IsPrimaryKey",
    "ForeignKeyReference",
    "DefaultValue",
    "ComputedDefinition",
    "Collation",
    "Description",
]

# Headers we insist on before we treat a row as the header row of the sheet.
_REQUIRED_HEADERS = {"tablename", "columnname"}


# ---------------------------------------------------------------------------
# Small value helpers
# ---------------------------------------------------------------------------
def _norm(value) -> str:
    """Return a trimmed string for any cell value (None -> "")."""
    if value is None:
        return ""
    return str(value).strip()


def _as_bool(value) -> bool:
    """Interpret common truthy spreadsheet conventions (1/Y/yes/true/x)."""
    text = _norm(value).lower()
    return text in {"1", "y", "yes", "true", "t", "x", "✓", "✔"}


def _header_key(value) -> str:
    """Normalise a header cell for matching (case/space/underscore-insensitive)."""
    return re.sub(r"[\s_]+", "", _norm(value)).lower()


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------
@dataclass
class Column:
    schema: str
    table: str
    order: int
    name: str
    data_type: str
    length: str = ""
    precision: str = ""
    scale: str = ""
    is_nullable: bool = True
    is_identity: bool = False
    is_computed: bool = False
    is_primary_key: bool = False
    foreign_key_reference: str = ""
    default_value: str = ""
    computed_definition: str = ""
    collation: str = ""
    description: str = ""

    def rendered_type(self) -> str:
        """A human-friendly type string, e.g. ``varchar(255)`` or ``decimal(18,2)``."""
        base = self.data_type or "unknown"
        length = _norm(self.length)
        precision = _norm(self.precision)
        scale = _norm(self.scale)
        if precision and precision not in ("0",):
            if scale and scale not in ("0",):
                return f"{base}({precision},{scale})"
            return f"{base}({precision})"
        if length and length not in ("0",):
            return f"{base}({length})"
        return base


@dataclass
class Table:
    schema: str
    name: str
    columns: list[Column] = field(default_factory=list)

    @property
    def full_name(self) -> str:
        return f"{self.schema}.{self.name}" if self.schema else self.name


@dataclass
class Relationship:
    parent_table: str  # referenced (the "one" side)
    child_table: str  # holds the foreign key (the "many" side)
    label: str = ""


@dataclass
class Schema:
    tables: list[Table] = field(default_factory=list)
    relationships: list[Relationship] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Reading the spreadsheet
# ---------------------------------------------------------------------------
def read_rows(path: str, progress: ProgressCallback | None = None) -> list[dict]:
    """Read an .xlsx/.xlsm file and return a list of ``{header: value}`` dicts.

    The header row is located by scanning for a row that contains the required
    headers, which lets the sheet have a title/banner above the real headers.

    ``progress``, if given, is called with a fraction (0.0 .. 1.0) as the rows
    are streamed in, so a caller can drive a progress bar for large files.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError(
            "openpyxl is required to read Excel files. Install it with "
            "'pip install openpyxl'."
        ) from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active

    # ``max_row`` is available for most files; when it isn't we simply can't
    # report a fraction, so the caller falls back to an indeterminate bar.
    total = sheet.max_row if isinstance(sheet.max_row, int) and sheet.max_row > 0 else 0
    grid = []
    for index, row in enumerate(sheet.iter_rows(values_only=True)):
        grid.append(list(row))
        # Report every so often to keep the UI responsive without flooding it.
        if progress is not None and total and index % 200 == 0:
            progress(min(index / total, 1.0))
    workbook.close()
    if progress is not None:
        progress(1.0)

    header_index = _find_header_row(grid)
    if header_index is None:
        raise ValueError(
            "Could not find a header row containing 'TableName' and 'ColumnName'. "
            "Make sure the sheet matches the expected schema layout."
        )

    headers = [_norm(cell) for cell in grid[header_index]]
    rows: list[dict] = []
    for raw in grid[header_index + 1 :]:
        record = {
            headers[i]: raw[i] if i < len(raw) else None
            for i in range(len(headers))
            if headers[i]
        }
        if any(_norm(v) for v in record.values()):
            rows.append(record)
    return rows


def _find_header_row(grid: list[list]) -> int | None:
    for index, row in enumerate(grid):
        keys = {_header_key(cell) for cell in row}
        if _REQUIRED_HEADERS.issubset(keys):
            return index
    return None


# ---------------------------------------------------------------------------
# Building the schema model from raw rows
# ---------------------------------------------------------------------------
def _build_getter(row: dict):
    """Return a lookup that matches headers loosely (case/space-insensitive)."""
    normalised = {_header_key(k): v for k, v in row.items()}

    def get(name: str):
        return normalised.get(_header_key(name))

    return get


def build_schema(
    rows: Iterable[dict], progress: ProgressCallback | None = None
) -> Schema:
    """Turn raw ``{header: value}`` rows into a :class:`Schema`.

    ``progress``, if given, is called with a fraction (0.0 .. 1.0) as rows are
    processed, so a caller can drive a progress bar for large schemas.
    """
    tables: dict[str, Table] = {}
    order_by_full: dict[str, Table] = {}

    rows = list(rows)
    total = len(rows)
    for index, row in enumerate(rows):
        if progress is not None and total and index % 200 == 0:
            progress(index / total)
        get = _build_getter(row)
        table_name = _norm(get("TableName"))
        column_name = _norm(get("ColumnName"))
        if not table_name or not column_name:
            continue

        schema_name = _norm(get("SchemaName"))
        key = f"{schema_name.lower()}::{table_name.lower()}"
        table = tables.get(key)
        if table is None:
            table = Table(schema=schema_name, name=table_name)
            tables[key] = table
            order_by_full[table_name.lower()] = table

        try:
            order = int(float(_norm(get("ColumnOrder"))))
        except (ValueError, TypeError):
            order = len(table.columns) + 1

        table.columns.append(
            Column(
                schema=schema_name,
                table=table_name,
                order=order,
                name=column_name,
                data_type=_norm(get("DataType")),
                length=_norm(get("Length")),
                precision=_norm(get("Precision")),
                scale=_norm(get("Scale")),
                is_nullable=_as_bool(get("IsNullable")),
                is_identity=_as_bool(get("IsIdentity")),
                is_computed=_as_bool(get("IsComputed")),
                is_primary_key=_as_bool(get("IsPrimaryKey")),
                foreign_key_reference=_norm(get("ForeignKeyReference")),
                default_value=_norm(get("DefaultValue")),
                computed_definition=_norm(get("ComputedDefinition")),
                collation=_norm(get("Collation")),
                description=_norm(get("Description")),
            )
        )

    if progress is not None:
        progress(1.0)

    ordered_tables = sorted(tables.values(), key=lambda t: (t.schema.lower(), t.name.lower()))
    for table in ordered_tables:
        table.columns.sort(key=lambda c: c.order)

    relationships = _derive_relationships(ordered_tables, order_by_full)
    return Schema(tables=ordered_tables, relationships=relationships)


def _parse_reference_table(reference: str) -> str | None:
    """Extract the referenced table name from a foreign-key reference string.

    Handles ``schema.Table.Column``, ``Table.Column``, ``Table(Column)`` and a
    bare ``Table``.
    """
    ref = _norm(reference)
    if not ref:
        return None
    # Strip a trailing "(Column)" part if present.
    ref = re.split(r"[(\[]", ref, maxsplit=1)[0].strip()
    if not ref:
        return None
    parts = [p for p in re.split(r"[.\s]+", ref) if p]
    if not parts:
        return None
    if len(parts) >= 3:
        return parts[-2]  # schema . TABLE . column
    if len(parts) == 2:
        return parts[0]  # TABLE . column
    return parts[0]  # bare table (or table with no column specified)


def _derive_relationships(
    tables: list[Table], by_name: dict[str, Table]
) -> list[Relationship]:
    seen: set[tuple[str, str, str]] = set()
    relationships: list[Relationship] = []
    for table in tables:
        for column in table.columns:
            if not column.foreign_key_reference:
                continue
            parent_name = _parse_reference_table(column.foreign_key_reference)
            if not parent_name:
                continue
            parent = by_name.get(parent_name.lower())
            parent_label = parent.name if parent else parent_name
            key = (parent_label.lower(), table.name.lower(), column.name.lower())
            if key in seen:
                continue
            seen.add(key)
            relationships.append(
                Relationship(
                    parent_table=parent_label,
                    child_table=table.name,
                    label=column.name,
                )
            )
    return relationships


# ---------------------------------------------------------------------------
# Mermaid generation
# ---------------------------------------------------------------------------
def _entity_id(name: str) -> str:
    """A Mermaid-safe entity identifier."""
    safe = re.sub(r"[^0-9A-Za-z_]", "_", name)
    if not safe:
        safe = "table"
    if safe[0].isdigit():
        safe = f"t_{safe}"
    return safe


def _attr_type(column: Column) -> str:
    """A Mermaid-safe attribute type token.

    Mermaid's ER parser only accepts a bare word for the attribute type, so
    ``varchar(100)`` and ``decimal(18,2)`` must be flattened to ``varchar_100``
    and ``decimal_18_2``.
    """
    token = re.sub(r"\s+", "", column.rendered_type())
    token = re.sub(r"[^0-9A-Za-z_]+", "_", token)  # ( ) , etc. -> underscore
    token = re.sub(r"_+", "_", token).strip("_")
    return token or "unknown"


def _attr_name(column: Column) -> str:
    return re.sub(r"[^0-9A-Za-z_]", "_", column.name) or "column"


def _quote_comment(text: str) -> str:
    """Sanitise a comment for the ``"..."`` slot of a Mermaid attribute."""
    return text.replace("\\", "/").replace('"', "'").replace("\n", " ").strip()


def generate_mermaid(schema: Schema) -> str:
    """Render a :class:`Schema` as a Mermaid ``erDiagram``."""
    lines: list[str] = ["erDiagram"]

    # Relationships first so the diagram reads top-down.
    for rel in schema.relationships:
        parent = _entity_id(rel.parent_table)
        child = _entity_id(rel.child_table)
        label = _quote_comment(rel.label) or "references"
        lines.append(f'    {parent} ||--o{{ {child} : "{label}"')

    if schema.relationships:
        lines.append("")

    for table in schema.tables:
        lines.append(f"    {_entity_id(table.name)} {{")
        for column in table.columns:
            keys = []
            if column.is_primary_key:
                keys.append("PK")
            if column.foreign_key_reference:
                keys.append("FK")
            key_part = f" {','.join(keys)}" if keys else ""

            notes = []
            if column.description:
                notes.append(column.description)
            if column.is_identity:
                notes.append("identity")
            if column.is_computed:
                notes.append("computed")
            if not column.is_nullable:
                notes.append("not null")
            comment = _quote_comment("; ".join(notes))
            comment_part = f' "{comment}"' if comment else ""

            lines.append(
                f"        {_attr_type(column)} {_attr_name(column)}{key_part}{comment_part}"
            )
        lines.append("    }")

    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# Convenience wrappers
# ---------------------------------------------------------------------------
def workbook_to_mermaid(path: str) -> tuple[Schema, str]:
    """Read an Excel file and return ``(schema, mermaid_text)``."""
    rows = read_rows(path)
    schema = build_schema(rows)
    return schema, generate_mermaid(schema)


_HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Database ER Diagram</title>
<style>
  :root {{ color-scheme: light dark; }}
  body {{ font-family: system-ui, -apple-system, Segoe UI, Roboto, sans-serif;
         margin: 0; padding: 24px; background: #ffffff; color: #1a1a1a; }}
  @media (prefers-color-scheme: dark) {{ body {{ background: #16181d; color: #e6e6e6; }} }}
  h1 {{ font-size: 18px; font-weight: 600; }}
  .diagram {{ overflow: auto; }}
</style>
<script src="https://cdn.jsdelivr.net/npm/mermaid@10/dist/mermaid.min.js"></script>
<script>
  mermaid.initialize({{ startOnLoad: true, securityLevel: 'loose', theme: 'default',
    maxTextSize: 2000000, maxEdges: 10000 }});
</script>
</head>
<body>
  <h1>Database ER Diagram</h1>
  <div class="diagram"><pre class="mermaid">
{diagram}
  </pre></div>
</body>
</html>
"""


def wrap_mermaid_html(mermaid_text: str) -> str:
    """Wrap Mermaid source in a standalone HTML page that renders it via CDN."""
    return _HTML_TEMPLATE.format(diagram=mermaid_text.strip())


if __name__ == "__main__":  # pragma: no cover
    import sys

    if len(sys.argv) < 2:
        print("Usage: python excel_to_mermaid.py <file.xlsx>")
        raise SystemExit(1)
    _schema, _mermaid = workbook_to_mermaid(sys.argv[1])
    print(_mermaid)
