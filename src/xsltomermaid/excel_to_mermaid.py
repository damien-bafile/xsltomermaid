"""Core logic: turn a database-schema spreadsheet into a Mermaid ER diagram.

The spreadsheet is expected to have one row per *column* of a database, with the
headers shown in the reference screenshot:

    SchemaName, TableName, ColumnOrder, ColumnName, DataType, Length, Precision,
    Scale, IsNullable, IsIdentity, IsComputed, IsPrimaryKey, ForeignKeyReference,
    DefaultValue, ComputedDefinition, Collation, Description

This module is deliberately free of any GUI or heavy dependencies for the parsing
and generation steps so it can be unit-tested on its own. ``openpyxl`` is only
imported inside :func:`read_rows` so the rest of the module works without it.
"""

from __future__ import annotations

import re
from html import escape as html_escape
from collections import defaultdict, deque
from dataclasses import dataclass, field
from functools import lru_cache
from typing import Callable, Iterable

# A progress callback receives a fraction in the range 0.0 .. 1.0.
ProgressCallback = Callable[[float], None]

# The canonical headers, in the order they appear in the reference sheet.
EXPECTED_HEADERS: list[str] = [
    "SchemaName",
    "TableName",
    "ColumnOrder",
    "ColumnName",
    "DataType",
    "Length",
    "Precision",
    "Scale",
    "IsNullable",
    "IsIdentity",
    "IsComputed",
    "IsPrimaryKey",
    "ForeignKeyReference",
    "DefaultValue",
    "ComputedDefinition",
    "Collation",
    "Description",
]

# Headers we insist on before we treat a row as the header row of the sheet.
_REQUIRED_HEADERS = {"tablename", "columnname"}


# ---------------------------------------------------------------------------
# Small value helpers
# ---------------------------------------------------------------------------
# Spreadsheet exports from SSMS/Dynamics/etc. commonly write the literal text
# "NULL" (or "(null)") into an empty cell rather than leaving it blank. Treat a
# cell whose *entire* value is one of these sentinels as empty, so an empty
# foreign-key cell doesn't look like a reference to a table called "NULL".
_NULL_TOKENS = {"null", "(null)"}


def _norm(value) -> str:
    """Return a trimmed string for any cell value (None -> "").

    A cell whose whole value is a null sentinel such as ``NULL`` is treated as
    empty; ``NOT NULL`` and other strings that merely contain the word are left
    untouched.
    """
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in _NULL_TOKENS:
        return ""
    return text


def _as_bool(value) -> bool:
    """Interpret common truthy spreadsheet conventions (1/Y/yes/true/x)."""
    text = _norm(value).lower()
    return text in {"1", "y", "yes", "true", "t", "x", "✓", "✔"}


_HEADER_KEY_RE = re.compile(r"[\s_]+")


@lru_cache(maxsize=4096)
def _header_key(value) -> str:
    """Normalise a header cell for matching (case/space/underscore-insensitive).

    Cached: headers and the expected field names are a tiny fixed set reused
    across every row, so build_schema was recomputing the same regex sub
    hundreds of thousands of times. Pure function of ``value`` (hashable cell
    scalars), so memoising is safe.
    """
    return _HEADER_KEY_RE.sub("", _norm(value)).lower()


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------
@dataclass
class Column:
    schema: str
    table: str
    order: int
    name: str
    data_type: str
    length: str = ""
    precision: str = ""
    scale: str = ""
    is_nullable: bool = True
    is_identity: bool = False
    is_computed: bool = False
    is_primary_key: bool = False
    foreign_key_reference: str = ""
    default_value: str = ""
    computed_definition: str = ""
    collation: str = ""
    description: str = ""

    def rendered_type(self) -> str:
        """A human-friendly type string, e.g. ``varchar(255)`` or ``decimal(18,2)``."""
        base = self.data_type or "unknown"
        length = _norm(self.length)
        precision = _norm(self.precision)
        scale = _norm(self.scale)
        if precision and precision not in ("0",):
            if scale and scale not in ("0",):
                return f"{base}({precision},{scale})"
            return f"{base}({precision})"
        if length and length not in ("0",):
            return f"{base}({length})"
        return base


@dataclass
class Table:
    schema: str
    name: str
    columns: list[Column] = field(default_factory=list)

    @property
    def full_name(self) -> str:
        return f"{self.schema}.{self.name}" if self.schema else self.name


@dataclass
class Relationship:
    parent_table: str  # referenced (the "one" side)
    child_table: str  # holds the foreign key (the "many" side)
    label: str = ""
    # The join, column for column: child_columns[i] references
    # parent_columns[i]. Empty, or of different lengths, when the sheet doesn't
    # say which parent column is referenced (e.g. a bare "Customer").
    child_columns: tuple[str, ...] = ()
    parent_columns: tuple[str, ...] = ()


@dataclass
class Schema:
    tables: list[Table] = field(default_factory=list)
    relationships: list[Relationship] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Reading the spreadsheet
# ---------------------------------------------------------------------------
def _coerce_cell(value):
    """Normalise a fast-reader cell to match openpyxl's typing.

    python-calamine returns whole numbers as floats (``1`` -> ``1.0``) and blank
    cells as ``""``; coerce those so downstream ``_norm`` / ``_as_bool`` /
    ``rendered_type`` behave identically to the openpyxl path (e.g. ``1.0`` would
    otherwise fail the ``_as_bool`` truthy check and drop PK/nullability flags).
    Booleans are left alone (``bool`` is a subclass of ``int``).
    """
    if isinstance(value, bool):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if value == "":
        return None
    return value


def _read_grid_openpyxl(
    path: str, progress: ProgressCallback | None = None
) -> list[list]:
    """Read the active sheet into a list of row lists using openpyxl."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError(
            "openpyxl is required to read Excel files. Install it with "
            "'pip install openpyxl'."
        ) from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active

    # ``max_row`` is available for most files; when it isn't we simply can't
    # report a fraction, so the caller falls back to an indeterminate bar.
    total = sheet.max_row if isinstance(sheet.max_row, int) and sheet.max_row > 0 else 0
    grid: list[list] = []
    for index, row in enumerate(sheet.iter_rows(values_only=True)):
        grid.append(list(row))
        # Report every so often to keep the UI responsive without flooding it.
        if progress is not None and total and index % 200 == 0:
            progress(min(index / total, 1.0))
    workbook.close()
    return grid


def _read_grid(path: str, progress: ProgressCallback | None = None) -> list[list]:
    """Read the sheet as a list of row lists, using the fastest reader available.

    Prefers python-calamine (a Rust-based xlsx reader, several times faster than
    openpyxl) and falls back to openpyxl if it isn't installed or can't read the
    file — so the app still works if the compiled wheel is unavailable. Both
    readers return equivalently-typed cells (see :func:`_coerce_cell`).
    """
    try:
        from python_calamine import CalamineWorkbook
    except ImportError:
        return _read_grid_openpyxl(path, progress)

    try:
        workbook = CalamineWorkbook.from_path(path)
        sheet = workbook.get_sheet_by_index(0)
        raw_rows = sheet.to_python()
    except Exception:  # noqa: BLE001 - any calamine hiccup -> proven openpyxl path
        return _read_grid_openpyxl(path, progress)

    # calamine parses the whole sheet in one call, so there's no per-row point to
    # stream a fraction from; it's fast enough that a single jump is fine.
    grid = [[_coerce_cell(cell) for cell in row] for row in raw_rows]
    if progress is not None:
        progress(1.0)
    return grid


def read_rows(path: str, progress: ProgressCallback | None = None) -> list[dict]:
    """Read an .xlsx/.xlsm file and return a list of ``{header: value}`` dicts.

    The header row is located by scanning for a row that contains the required
    headers, which lets the sheet have a title/banner above the real headers.

    ``progress``, if given, is called with a fraction (0.0 .. 1.0) as the file is
    read, so a caller can drive a progress bar for large files.
    """
    grid = _read_grid(path, progress)
    if progress is not None:
        progress(1.0)

    header_index = _find_header_row(grid)
    if header_index is None:
        raise ValueError(
            "Could not find a header row containing 'TableName' and 'ColumnName'. "
            "Make sure the sheet matches the expected schema layout."
        )

    headers = [_norm(cell) for cell in grid[header_index]]
    rows: list[dict] = []
    for raw in grid[header_index + 1 :]:
        record = {
            headers[i]: raw[i] if i < len(raw) else None
            for i in range(len(headers))
            if headers[i]
        }
        if any(_norm(v) for v in record.values()):
            rows.append(record)
    return rows


def _find_header_row(grid: list[list]) -> int | None:
    for index, row in enumerate(grid):
        keys = {_header_key(cell) for cell in row}
        if _REQUIRED_HEADERS.issubset(keys):
            return index
    return None


# ---------------------------------------------------------------------------
# Building the schema model from raw rows
# ---------------------------------------------------------------------------
def _build_getter(row: dict):
    """Return a lookup that matches headers loosely (case/space-insensitive)."""
    normalised = {_header_key(k): v for k, v in row.items()}

    def get(name: str):
        return normalised.get(_header_key(name))

    return get


def build_schema(
    rows: Iterable[dict], progress: ProgressCallback | None = None
) -> Schema:
    """Turn raw ``{header: value}`` rows into a :class:`Schema`.

    ``progress``, if given, is called with a fraction (0.0 .. 1.0) as rows are
    processed, so a caller can drive a progress bar for large schemas.
    """
    tables: dict[str, Table] = {}
    order_by_full: dict[str, Table] = {}

    rows = list(rows)
    total = len(rows)
    for index, row in enumerate(rows):
        if progress is not None and total and index % 200 == 0:
            progress(index / total)
        get = _build_getter(row)
        table_name = _norm(get("TableName"))
        column_name = _norm(get("ColumnName"))
        if not table_name or not column_name:
            continue

        schema_name = _norm(get("SchemaName"))
        key = f"{schema_name.lower()}::{table_name.lower()}"
        table = tables.get(key)
        if table is None:
            table = Table(schema=schema_name, name=table_name)
            tables[key] = table
            order_by_full[table_name.lower()] = table

        try:
            order = int(float(_norm(get("ColumnOrder"))))
        except (ValueError, TypeError):
            order = len(table.columns) + 1

        table.columns.append(
            Column(
                schema=schema_name,
                table=table_name,
                order=order,
                name=column_name,
                data_type=_norm(get("DataType")),
                length=_norm(get("Length")),
                precision=_norm(get("Precision")),
                scale=_norm(get("Scale")),
                is_nullable=_as_bool(get("IsNullable")),
                is_identity=_as_bool(get("IsIdentity")),
                is_computed=_as_bool(get("IsComputed")),
                is_primary_key=_as_bool(get("IsPrimaryKey")),
                foreign_key_reference=_norm(get("ForeignKeyReference")),
                default_value=_norm(get("DefaultValue")),
                computed_definition=_norm(get("ComputedDefinition")),
                collation=_norm(get("Collation")),
                description=_norm(get("Description")),
            )
        )

    if progress is not None:
        progress(1.0)

    ordered_tables = sorted(tables.values(), key=lambda t: (t.schema.lower(), t.name.lower()))
    for table in ordered_tables:
        table.columns.sort(key=lambda c: c.order)
        _drop_solution_layering_keys(table)

    relationships = _derive_relationships(ordered_tables, order_by_full)
    return Schema(tables=ordered_tables, relationships=relationships)


# Dynamics 365 / Dataverse solution-layering columns. Their physical primary key
# is (id, overwritetime, componentstate) so each solution layer of a component
# can be stored; to a reader of the diagram (and to every foreign key) the key
# is just the id.
SOLUTION_LAYERING_KEYS = frozenset({"overwritetime", "componentstate"})


def _drop_solution_layering_keys(table: Table) -> None:
    """Stop marking solution-layering columns as PK, when a real key remains."""
    keys = [c for c in table.columns if c.is_primary_key]
    layering = [c for c in keys if c.name.lower() in SOLUTION_LAYERING_KEYS]
    if layering and len(layering) < len(keys):
        for column in layering:
            column.is_primary_key = False


_REF_COLUMNS_RE =re.compile(r"^(?P<table>[^(]*)\((?P<columns>[^)]*)\)\s*$")


def _parse_reference(reference: str) -> tuple[str | None, list[str]]:
    """Split a foreign-key reference into ``(table, [referenced columns])``.

    Handles ``schema.Table.Column``, ``Table.Column``, ``Table(Column)``,
    ``schema.Table(ColA, ColB)`` (a composite key in one cell), a bare
    ``Table``, and SQL-Server-style ``[schema].[Table].[Column]`` quoting.
    """
    ref = _norm(reference)
    if not ref:
        return None, []
    match = _REF_COLUMNS_RE.match(ref)
    if match:
        # "(Column[, Column])" names the columns, so the rest is [schema.]Table.
        columns = [
            c for c in (re.sub(r'[\[\]"`]', "", p).strip()
                        for p in match["columns"].split(","))
            if c
        ]
        ref = match["table"]
    else:
        columns = []
    parts = [p for p in re.split(r"[.\s]+", re.sub(r'[\[\]"`]', "", ref)) if p]
    if not parts:
        return None, columns
    if match:
        return parts[-1], columns  # [schema .] TABLE ( columns )
    if len(parts) >= 3:
        return parts[-2], [parts[-1]]  # schema . TABLE . column
    if len(parts) == 2:
        return parts[0], [parts[1]]  # TABLE . column
    return parts[0], []  # bare table (or table with no column specified)


def _parse_reference_table(reference: str) -> str | None:
    """Extract the referenced table name from a foreign-key reference string."""
    return _parse_reference(reference)[0]


def _derive_relationships(
    tables: list[Table], by_name: dict[str, Table]
) -> list[Relationship]:
    """One relationship per foreign key.

    A composite foreign key arrives as several FK columns in the child, each
    referencing one column of the parent's composite primary key. Those are
    merged into a single relationship (labelled ``ColA, ColB``) rather than
    drawing one parallel edge per column. Separate single-column FKs to the
    same parent (e.g. ``createdby`` / ``modifiedby``) stay separate edges.
    """
    seen: set[tuple[str, str, str]] = set()
    relationships: list[Relationship] = []

    def add(
        parent_label: str, child: str, columns: list[Column], refs: list[str]
    ) -> None:
        label = ", ".join(c.name for c in columns)
        key = (parent_label.lower(), child.lower(), label.lower())
        if key in seen:
            return
        seen.add(key)
        relationships.append(
            Relationship(
                parent_table=parent_label,
                child_table=child,
                label=label,
                child_columns=tuple(c.name for c in columns),
                parent_columns=tuple(refs),
            )
        )

    def canonical(parent: Table | None, name: str) -> str:
        """The parent's own spelling of a referenced column, when it has one."""
        if parent is not None:
            for c in parent.columns:
                if c.name.lower() == name.lower():
                    return c.name
        return name

    for table in tables:
        # Composite-key candidates, per parent: FK columns in column order
        # that reference one of the parent's (multi-column) primary key.
        pending: dict[str, list[tuple[Column, str]]] = {}
        labels: dict[str, str] = {}

        parents: dict[str, Table | None] = {}

        def flush(parent_key: str) -> None:
            parent = parents[parent_key]
            for column, ref in pending.pop(parent_key, []):
                add(labels[parent_key], table.name, [column], [canonical(parent, ref)])

        for column in table.columns:
            if not column.foreign_key_reference:
                continue
            parent_name, ref_columns = _parse_reference(column.foreign_key_reference)
            if not parent_name:
                continue
            parent = by_name.get(parent_name.lower())
            parent_label = parent.name if parent else parent_name
            parent_key = parent_label.lower()
            pk = {c.name.lower() for c in parent.columns if c.is_primary_key} if parent else set()
            ref = ref_columns[0].lower() if len(ref_columns) == 1 else ""

            if len(pk) < 2 or ref not in pk:
                if ref_columns:
                    refs = [canonical(parent, r) for r in ref_columns]
                else:
                    # A bare "Table" reference means its primary key — when
                    # that key is a single column, the join is unambiguous.
                    keys = [c.name for c in parent.columns if c.is_primary_key] if parent else []
                    refs = keys if len(keys) == 1 else []
                add(parent_label, table.name, [column], refs)
                continue

            labels[parent_key] = parent_label
            parents[parent_key] = parent
            group = pending.setdefault(parent_key, [])
            if any(r == ref for _, r in group):
                # The same key column again: a second FK has started, so the
                # one in progress was partial — keep its columns as-is.
                flush(parent_key)
                group = pending.setdefault(parent_key, [])
            group.append((column, ref))
            if {r for _, r in group} == pk:
                done = pending.pop(parent_key)
                add(
                    parent_label,
                    table.name,
                    [c for c, _ in done],
                    [canonical(parent, r) for _, r in done],
                )

        for parent_key in list(pending):
            flush(parent_key)
    return relationships


def filter_schema(
    schema: Schema,
    selected_names: Iterable[str],
    include_related: bool = False,
) -> Schema:
    """Return a copy of ``schema`` containing only the selected tables.

    ``selected_names`` is matched against table names case-insensitively.
    Relationships are kept only when *both* endpoints are in the result, so the
    filtered diagram never dangles an edge to a table that isn't drawn.

    When ``include_related`` is true, the *direct* foreign-key neighbours of the
    selected tables (in either direction) are pulled in as well — a single layer
    out, so a picked table is shown with what it references and what references
    it, but not those tables' further neighbours.
    """
    wanted = {name.strip().lower() for name in selected_names if name.strip()}

    if include_related and wanted:
        # Match against the frozen original selection so neighbours-of-neighbours
        # aren't dragged in; this expands by exactly one hop.
        selected = set(wanted)
        for rel in schema.relationships:
            parent = rel.parent_table.lower()
            child = rel.child_table.lower()
            if parent in selected or child in selected:
                wanted.add(parent)
                wanted.add(child)

    tables = [t for t in schema.tables if t.name.lower() in wanted]
    present = {t.name.lower() for t in tables}
    relationships = [
        rel
        for rel in schema.relationships
        if rel.parent_table.lower() in present and rel.child_table.lower() in present
    ]
    return Schema(tables=tables, relationships=relationships)


def filter_columns(
    schema: Schema, excluded: set[tuple[str, str]]
) -> Schema:
    """Return a copy of ``schema`` with specific columns dropped.

    ``excluded`` is a set of ``(table_name_lower, column_name_lower)`` pairs; any
    column matching one is left out of its table. Relationships are kept as-is
    (they connect tables, not individual columns), so hiding a foreign-key
    column doesn't remove the edge it implied.
    """
    if not excluded:
        return schema

    new_tables = []
    for table in schema.tables:
        kept = [
            column
            for column in table.columns
            if (table.name.lower(), column.name.lower()) not in excluded
        ]
        new_tables.append(Table(schema=table.schema, name=table.name, columns=kept))
    return Schema(tables=new_tables, relationships=schema.relationships)


def shortest_path(schema: Schema, start: str, end: str) -> list[str] | None:
    """Shortest chain of tables linking ``start`` to ``end`` over the FK graph.

    Relationships are treated as undirected edges. Returns the list of table
    names on a shortest path, including both endpoints (canonical casing), or
    ``None`` if the two tables aren't connected. Names are matched
    case-insensitively; an endpoint that isn't a table in the schema yields
    ``None``.
    """
    canon = {t.name.lower(): t.name for t in schema.tables}
    src = start.strip().lower()
    dst = end.strip().lower()
    if src not in canon or dst not in canon:
        return None
    if src == dst:
        return [canon[src]]

    adjacency: dict[str, set[str]] = {}
    for rel in schema.relationships:
        parent = rel.parent_table.lower()
        child = rel.child_table.lower()
        adjacency.setdefault(parent, set()).add(child)
        adjacency.setdefault(child, set()).add(parent)

    # Breadth-first search records each node's predecessor for reconstruction.
    previous: dict[str, str | None] = {src: None}
    queue: deque[str] = deque([src])
    while queue:
        node = queue.popleft()
        if node == dst:
            break
        for neighbour in adjacency.get(node, ()):
            if neighbour not in previous:
                previous[neighbour] = node
                queue.append(neighbour)

    if dst not in previous:
        return None

    path: list[str] = []
    node: str | None = dst
    while node is not None:
        path.append(canon.get(node, node))
        node = previous[node]
    path.reverse()
    return path


# Foreign keys are directed (a child table references a parent), so a "path"
# can mean three different things. ``either`` ignores direction (mere
# connectivity); ``forward`` follows the FK reference child→parent ("what does
# this depend on"); ``reverse`` follows it parent→child ("what depends on this").
PATH_DIRECTIONS = ("either", "forward", "reverse")


def _path_adjacency(schema: Schema, direction: str) -> dict[str, set[str]]:
    """Build the (possibly directed) table adjacency for path finding."""
    adjacency: dict[str, set[str]] = defaultdict(set)
    for rel in schema.relationships:
        parent = rel.parent_table.lower()
        child = rel.child_table.lower()
        if direction in ("either", "forward"):
            adjacency[child].add(parent)  # follow the FK reference
        if direction in ("either", "reverse"):
            adjacency[parent].add(child)  # follow the dependents
    return adjacency


def all_shortest_paths(
    schema: Schema,
    start: str,
    end: str,
    direction: str = "either",
    limit: int = 12,
) -> list[list[str]]:
    """Every shortest chain of tables from ``start`` to ``end`` over the FK graph.

    Like :func:`shortest_path` but returns *all* equally-short paths (so the UI
    can let the user choose when there's a tie), honouring ``direction`` (one of
    :data:`PATH_DIRECTIONS`). Results are sorted for determinism and capped at
    ``limit``. Returns ``[]`` when the two tables aren't connected under the
    chosen direction, or ``[[start]]`` when they're the same table.
    """
    canon = {t.name.lower(): t.name for t in schema.tables}
    src = start.strip().lower()
    dst = end.strip().lower()
    if src not in canon or dst not in canon:
        return []
    if src == dst:
        return [[canon[src]]]

    adjacency = _path_adjacency(schema, direction)
    # BFS layers: distance from src to every reachable node.
    dist: dict[str, int] = {src: 0}
    queue: deque[str] = deque([src])
    while queue:
        node = queue.popleft()
        for neighbour in adjacency.get(node, ()):
            if neighbour not in dist:
                dist[neighbour] = dist[node] + 1
                queue.append(neighbour)
    if dst not in dist:
        return []

    # Predecessors on some shortest path: an edge n→m with dist[m] == dist[n]+1.
    preds: dict[str, list[str]] = defaultdict(list)
    for node, d in dist.items():
        for neighbour in adjacency.get(node, ()):
            if dist.get(neighbour) == d + 1:
                preds[neighbour].append(node)

    results: list[list[str]] = []

    def backtrack(node: str, tail: list[str]) -> None:
        if len(results) >= limit:
            return
        if node == src:
            results.append([canon.get(name, name) for name in [node] + tail])
            return
        for prev in sorted(preds[node]):
            backtrack(prev, [node] + tail)

    backtrack(dst, [])
    results.sort()
    return results[:limit]


def route_paths(
    schema: Schema,
    stops: list[str],
    direction: str = "either",
    limit: int = 12,
) -> tuple[list[list[str]], tuple[str, str] | None]:
    """Shortest routes visiting ``stops`` in order (a multi-waypoint path).

    Returns ``(routes, broken)``. ``broken`` is the first consecutive
    ``(from, to)`` pair with no connecting path (and ``routes`` is then empty);
    otherwise it's ``None`` and ``routes`` holds up to ``limit`` full routes,
    each stitched from a shortest path per segment.
    """
    clean: list[str] = []
    for stop in stops:
        name = (stop or "").strip()
        if name and (not clean or clean[-1].lower() != name.lower()):
            clean.append(name)
    if len(clean) < 2:
        return ([clean] if clean else []), None

    segments: list[list[list[str]]] = []
    for a, b in zip(clean, clean[1:]):
        options = all_shortest_paths(schema, a, b, direction, limit)
        if not options:
            return [], (a, b)
        segments.append(options)

    # Stitch one option per segment, joining on the shared endpoint. Capped so a
    # many-waypoint route with several options each can't explode the list.
    routes: list[list[str]] = [[]]
    for options in segments:
        combined: list[list[str]] = []
        for base in routes:
            for path in options:
                combined.append(base + (path if not base else path[1:]))
                if len(combined) >= limit:
                    break
            if len(combined) >= limit:
                break
        routes = combined
    return routes[:limit], None


def related_tables(
    schema: Schema,
    selected_names: Iterable[str],
    direction: str = "either",
) -> set[str]:
    """The one-hop foreign-key neighbours of ``selected_names`` (canonical names).

    ``direction`` (one of :data:`PATH_DIRECTIONS`) chooses which links to follow:
    ``either`` for both, ``forward`` for the tables the selection *references*
    (its parents), ``reverse`` for the tables that *reference* the selection
    (its children). The selected tables themselves are never returned.
    """
    canon = {t.name.lower(): t.name for t in schema.tables}
    selected = {name.strip().lower() for name in selected_names if name.strip()}
    adjacency = _path_adjacency(schema, direction)
    neighbours: set[str] = set()
    for name in selected:
        neighbours |= adjacency.get(name, set())
    return {canon[n] for n in neighbours if n not in selected and n in canon}


def unresolved_foreign_keys(schema: Schema) -> list[tuple[str, str]]:
    """Foreign-key references that don't resolve to a table in ``schema``.

    Returns ``(child_table, reference_string)`` pairs — a dangling FK points at a
    table that isn't in the sheet, so no relationship is drawn for it. Useful for
    telling the user their export may be missing tables.
    """
    names = {t.name.lower() for t in schema.tables}
    dangling: list[tuple[str, str]] = []
    for table in schema.tables:
        for column in table.columns:
            ref = _norm(column.foreign_key_reference)
            if not ref:
                continue
            target = _parse_reference_table(ref)
            if target is None or target.lower() not in names:
                dangling.append((table.name, ref))
    return dangling


# ---------------------------------------------------------------------------
# Mermaid generation
# ---------------------------------------------------------------------------
def _entity_id(name: str) -> str:
    """A Mermaid-safe entity identifier."""
    safe = re.sub(r"[^0-9A-Za-z_]", "_", name)
    if not safe:
        safe = "table"
    if safe[0].isdigit():
        safe = f"t_{safe}"
    return safe


def _attr_type(column: Column) -> str:
    """A Mermaid-safe attribute type token.

    Mermaid's ER parser only accepts a bare word for the attribute type, so
    ``varchar(100)`` and ``decimal(18,2)`` must be flattened to ``varchar_100``
    and ``decimal_18_2``.
    """
    token = re.sub(r"\s+", "", column.rendered_type())
    token = re.sub(r"[^0-9A-Za-z_]+", "_", token)  # ( ) , etc. -> underscore
    token = re.sub(r"_+", "_", token).strip("_")
    return token or "unknown"


def _attr_name(column: Column) -> str:
    return re.sub(r"[^0-9A-Za-z_]", "_", column.name) or "column"


def _quote_comment(text: str) -> str:
    """Sanitise a comment for the ``"..."`` slot of a Mermaid attribute."""
    return text.replace("\\", "/").replace('"', "'").replace("\n", " ").strip()


@dataclass
class DiagramOptions:
    """Content toggles that change the generated Mermaid *text*.

    These require regenerating the source (unlike :class:`RenderStyle`, which is
    applied at render time). Defaults reproduce the original output exactly.
    """

    show_comments: bool = True  # the "..." note slot (description, identity, …)
    show_rel_labels: bool = True  # the FK column name on a relationship line
    prefix_schema: bool = False  # entity id as schema.Table vs Table
    keys_only: bool = False  # show only PK/FK attributes


def mermaid_entity_ids(schema: Schema, options: DiagramOptions | None = None) -> dict[str, str]:
    """``{entity id: table name}`` for the ids :func:`generate_mermaid` writes.

    Lets the GUI map a click on a drawn table back to the table it shows.
    """
    opts = options or DiagramOptions()
    return {
        _entity_id(t.full_name if opts.prefix_schema else t.name): t.name
        for t in schema.tables
    }


def generate_mermaid(schema: Schema, options: DiagramOptions | None = None) -> str:
    """Render a :class:`Schema` as a Mermaid ``erDiagram``."""
    opts = options or DiagramOptions()
    lines: list[str] = ["erDiagram"]

    # Map table names to their Mermaid ids, so relationships and entity blocks
    # agree even when the schema prefix is switched on.
    by_name = {t.name.lower(): t for t in schema.tables}

    def entity_id(name: str) -> str:
        table = by_name.get(name.lower())
        label = table.full_name if (table and opts.prefix_schema) else name
        return _entity_id(label)

    # Relationships first so the diagram reads top-down.
    for rel in schema.relationships:
        parent = entity_id(rel.parent_table)
        child = entity_id(rel.child_table)
        if opts.show_rel_labels:
            label = _quote_comment(rel.label) or "references"
        else:
            label = ""
        lines.append(f'    {parent} ||--o{{ {child} : "{label}"')

    if schema.relationships:
        lines.append("")

    for table in schema.tables:
        lines.append(f"    {entity_id(table.name)} {{")
        for column in table.columns:
            if opts.keys_only and not (
                column.is_primary_key or column.foreign_key_reference
            ):
                continue

            keys = []
            if column.is_primary_key:
                keys.append("PK")
            if column.foreign_key_reference:
                keys.append("FK")
            key_part = f" {','.join(keys)}" if keys else ""

            comment_part = ""
            if opts.show_comments:
                notes = []
                if column.description:
                    notes.append(column.description)
                if column.is_identity:
                    notes.append("identity")
                if column.is_computed:
                    notes.append("computed")
                if not column.is_nullable:
                    notes.append("not null")
                comment = _quote_comment("; ".join(notes))
                comment_part = f' "{comment}"' if comment else ""

            lines.append(
                f"        {_attr_type(column)} {_attr_name(column)}{key_part}{comment_part}"
            )
        lines.append("    }")

    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# Convenience wrappers
# ---------------------------------------------------------------------------
def workbook_to_mermaid(path: str) -> tuple[Schema, str]:
    """Read an Excel file and return ``(schema, mermaid_text)``."""
    rows = read_rows(path)
    schema = build_schema(rows)
    return schema, generate_mermaid(schema)


_HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Database ER Diagram</title>
<style>
  :root {{ color-scheme: light dark; }}
  body {{ font-family: system-ui, -apple-system, Segoe UI, Roboto, sans-serif;
         margin: 0; padding: 24px; background: #ffffff; color: #1a1a1a; }}
  @media (prefers-color-scheme: dark) {{ body {{ background: #16181d; color: #e6e6e6; }} }}
  h1 {{ font-size: 18px; font-weight: 600; }}
  .diagram {{ overflow: auto; }}
</style>
<script src="{script_src}"></script>
<script>
  mermaid.initialize({{ startOnLoad: true, securityLevel: 'loose', theme: 'default',
    maxTextSize: 2000000, maxEdges: 10000 }});
</script>
</head>
<body>
  <h1>Database ER Diagram</h1>
  <div class="diagram"><pre class="mermaid">
{diagram}
  </pre></div>
</body>
</html>
"""


MERMAID_CDN = "https://cdn.jsdelivr.net/npm/mermaid@10/dist/mermaid.min.js"


def wrap_mermaid_html(mermaid_text: str, script_src: str = MERMAID_CDN) -> str:
    """Wrap Mermaid source in a standalone HTML page that renders it.

    ``script_src`` is where the page loads mermaid.js from: the CDN by default,
    or a local path (e.g. ``"mermaid.min.js"`` beside the page) to work offline.
    The source is HTML-escaped; Mermaid reads the element's text, so names
    containing ``<`` or ``&`` survive intact.
    """
    return _HTML_TEMPLATE.format(
        diagram=html_escape(mermaid_text.strip()),
        script_src=html_escape(script_src, quote=True),
    )


if __name__ == "__main__":  # pragma: no cover
    import sys

    if len(sys.argv) < 2:
        print("Usage: python -m xsltomermaid.excel_to_mermaid <file.xlsx>")
        raise SystemExit(1)
    _schema, _mermaid = workbook_to_mermaid(sys.argv[1])
    print(_mermaid)
