"""Headless GUI test: build the window, load the sample, and screenshot it.

Runs Qt with the offscreen platform so it works in CI with no display. The test
asserts the window actually painted something (a PNG of a sensible size).

Run: QT_QPA_PLATFORM=offscreen python test_screenshot.py
 or: uv run --extra – ...  (see README); pytest also picks it up.
"""

from __future__ import annotations

import os
import subprocess
import sys

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
# Needed for the WebEngine (Chromium) diagram render to run headless / as root.
os.environ.setdefault("QTWEBENGINE_DISABLE_SANDBOX", "1")
os.environ.setdefault(
    "QTWEBENGINE_CHROMIUM_FLAGS", "--no-sandbox --disable-gpu --in-process-gpu"
)

import struct
import tempfile
import xml.etree.ElementTree as ET

import pytest

# Skip cleanly if the GUI stack (PySide6 + system libs) isn't importable.
pytest.importorskip("PySide6")

from PySide6.QtWidgets import QApplication  # noqa: E402

import main as app_module  # noqa: E402
from make_sample import ROWS  # noqa: E402


def test_fit_scale_caps_large_exports():
    """A huge diagram at high scale is clamped to stay within the size caps."""
    from diagram_view import MAX_PNG_DIM, MAX_PNG_PIXELS, _fit_scale

    # Small diagram: requested scale is kept as-is.
    assert _fit_scale(400, 300, 2.0) == 2.0

    # Large diagram at 4x would be ~768 MP; must clamp under both caps.
    eff = _fit_scale(8000, 6000, 4.0)
    assert eff < 4.0
    assert 8000 * eff <= MAX_PNG_DIM + 1
    assert 6000 * eff <= MAX_PNG_DIM + 1
    assert (8000 * eff) * (6000 * eff) <= MAX_PNG_PIXELS * 1.01

    # Degenerate sizes don't blow up.
    assert _fit_scale(0, 0, 3.0) == 3.0


def test_render_style_defaults_to_left_to_right():
    from diagram_view import RenderStyle

    assert RenderStyle().layout_direction == "LR"


def test_diagram_options_bar_constructs_headlessly():
    app = QApplication.instance() or QApplication([])
    bar = app_module.DiagramOptionsBar()

    assert bar.render_style().layout_direction == "LR"


def test_column_selector_marks_primary_and_foreign_keys():
    from excel_to_mermaid import Column, Table

    app = QApplication.instance() or QApplication([])
    selector = app_module.ColumnSelector()
    selector.set_tables(
        [
            Table(
                "dbo",
                "Orders",
                [
                    Column("dbo", "Orders", 1, "OrderID", "int", is_primary_key=True),
                    Column(
                        "dbo",
                        "Orders",
                        2,
                        "CustomerID",
                        "int",
                        foreign_key_reference="dbo.Customer.CustomerID",
                    ),
                ],
            )
        ]
    )

    table_item = selector._tree.topLevelItem(0)
    assert table_item.child(0).text(0) == "OrderID  [PK]"
    assert table_item.child(1).text(0) == "CustomerID  [FK]"
    assert table_item.child(0).toolTip(0) == "PK key column: OrderID"
    del app


def test_sql_query_dock_is_fixed_on_right_and_copyable():
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    window.show()
    app.processEvents()
    assert window._sql_dock.isHidden()
    assert not window._sql_dock.toggleViewAction().isChecked()
    window._sql_dock.toggleViewAction().trigger()
    assert window._sql_dock.isVisible()

    assert window.dockWidgetArea(window._sql_dock) == app_module.Qt.RightDockWidgetArea
    assert window._sql_dock.allowedAreas() == app_module.Qt.RightDockWidgetArea
    assert not window._sql_dock.isFloating()
    assert window._sql_dock.features() == app_module.QDockWidget.DockWidgetClosable
    for header in app_module.EXPECTED_HEADERS:
        assert f"[{header}]" in window._sql_query_view.toPlainText()

    window.copy_sql_query()
    assert app.clipboard().text() == app_module.SQL_SERVER_SCHEMA_QUERY
    window._diagram_view.cleanup()
    window.close()
    del app


@pytest.mark.parametrize(
    "mode,pk,fk,expected",
    [
        ("order", False, False, ["zeta", "Id", "ParentId", "alpha"]),
        ("name", False, False, ["alpha", "Id", "ParentId", "zeta"]),
        ("name_desc", False, False, ["zeta", "ParentId", "Id", "alpha"]),
        ("type", False, False, ["alpha", "Id", "ParentId", "zeta"]),
        ("name", True, False, ["Id", "alpha", "ParentId", "zeta"]),
        ("name", False, True, ["ParentId", "alpha", "Id", "zeta"]),
        ("name_desc", True, True, ["Id", "ParentId", "zeta", "alpha"]),
    ],
)
def test_column_sorting_preserves_selection_and_source(mode, pk, fk, expected):
    from excel_to_mermaid import Column, Schema, Table, filter_columns, generate_mermaid

    app = QApplication.instance() or QApplication([])
    table = Table("dbo", "Example", [
        Column("dbo", "Example", 1, "zeta", "varchar"),
        Column("dbo", "Example", 2, "Id", "int", is_primary_key=True),
        Column("dbo", "Example", 3, "ParentId", "int",
               foreign_key_reference="dbo.Parent.Id"),
        Column("dbo", "Example", 4, "alpha", "bit"),
    ])
    selector = app_module.ColumnSelector()
    selector.set_tables([table])
    selector._tree.topLevelItem(0).child(0).setCheckState(0, app_module.Qt.Unchecked)
    selector._filter.setText("Id")
    selector._sort.setCurrentIndex(selector._sort.findData(mode))
    selector._pk_first.setChecked(pk)
    selector._fk_first.setChecked(fk)

    assert [c.name for c in selector.ordered_columns(table)] == expected
    children = list(selector._children(selector._tree.topLevelItem(0)))
    assert [c.text(0).split("  [")[0] for c in children] == expected
    assert next(c for c in children if c.text(0) == "zeta").checkState(0) == app_module.Qt.Unchecked
    assert next(c for c in children if c.text(0) == "alpha").isHidden()
    assert selector.excluded_pairs() == {("example", "zeta")}
    schema = selector.sorted_schema(filter_columns(Schema([table], []), selector.excluded_pairs()))
    assert [c.name for c in schema.tables[0].columns] == [n for n in expected if n != "zeta"]
    mermaid = generate_mermaid(schema)
    positions = [mermaid.index(f" {name}") for name in expected if name != "zeta"]
    assert positions == sorted(positions)
    assert [c.name for c in table.columns] == ["zeta", "Id", "ParentId", "alpha"]
    del app


def test_extracted_data_header_sorting_is_numeric_and_non_destructive():
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    rows = [
        [str(order) if h == "ColumnOrder" else name if h == "ColumnName" else ""
         for h in app_module.EXPECTED_HEADERS]
        for order, name in [(10, "zeta"), (2, "Alpha"), (1, "beta")]
    ]
    window._table_model.set_rows(rows)
    order_column = app_module.EXPECTED_HEADERS.index("ColumnOrder")
    name_column = app_module.EXPECTED_HEADERS.index("ColumnName")
    proxy = window._table.model()
    window._table.sortByColumn(order_column, app_module.Qt.AscendingOrder)
    assert [proxy.index(i, order_column).data() for i in range(3)] == ["1", "2", "10"]
    window._table.sortByColumn(order_column, app_module.Qt.DescendingOrder)
    assert [proxy.index(i, order_column).data() for i in range(3)] == ["10", "2", "1"]
    window._table.sortByColumn(name_column, app_module.Qt.AscendingOrder)
    assert [proxy.index(i, name_column).data() for i in range(3)] == ["Alpha", "beta", "zeta"]
    assert window._table_model._rows == rows
    window._diagram_view.cleanup()
    del app


def test_diagram_zoom_label_tracks_webengine_zoom():
    import diagram_view
    from PySide6.QtTest import QTest

    if not diagram_view.WEBENGINE_AVAILABLE:
        pytest.skip("PySide6 WebEngine not available")
    app = QApplication.instance() or QApplication([])
    view = diagram_view.DiagramView()
    assert view._zoom_label.text() == "Zoom: 100%"
    for zoom, text in [(1.25, "125%"), (0.5, "50%"), (1.0, "100%")]:
        view._view.setZoomFactor(zoom)
        QTest.qWait(250)
        assert view._zoom_label.text() == f"Zoom: {text}"
    view.cleanup()
    assert not view._zoom_timer.isActive()
    del app


def test_schema_to_drawio_creates_table_vertices_and_edges():
    import diagram_view
    from excel_to_mermaid import build_schema

    records = [dict(zip(app_module.EXPECTED_HEADERS, row)) for row in ROWS]
    schema = build_schema(records)
    drawio = diagram_view.schema_to_drawio(schema)
    root = ET.fromstring(drawio)

    assert root.tag == "mxfile"
    vertices = root.findall(".//mxCell[@vertex='1']")
    edges = root.findall(".//mxCell[@edge='1']")
    assert len(vertices) == len(schema.tables)
    assert len(edges) == len(schema.relationships)
    assert all("shape=mxgraph.er.entity" in cell.attrib.get("style", "") for cell in vertices)
    assert all("endArrow=ERmany" in cell.attrib.get("style", "") for cell in edges)
    # Every edge must carry a geometry, or Draw.io opens the saved file blank
    # (the import path is lenient, but File > Open is not).
    assert edges, "expected at least one relationship edge"
    assert all(cell.find("mxGeometry") is not None for cell in edges)


def test_svg_dimensions_handles_fractional_and_exponent_sizes():
    import diagram_view

    w, h = diagram_view._svg_dimensions('<svg width=".5" height="1e3"></svg>')
    assert w == 0.5
    assert h == 1000.0
    w1, h1 = diagram_view._svg_dimensions('<svg width="1e3px" height="2.5e2pt"></svg>')
    assert w1 == 1000.0
    assert h1 == 250.0
    w0, h0 = diagram_view._svg_dimensions('<svg width="0" height="0"></svg>')
    assert w0 == 0.0
    assert h0 == 0.0
    # A percentage width (Mermaid's useMaxWidth) must fall back to the viewBox,
    # not be read as 100 — otherwise the PDF fallback would size to 100px.
    wp, hp = diagram_view._svg_dimensions(
        '<svg width="100%" viewBox="0 0 3526 6531"></svg>'
    )
    assert wp == 3526.0
    assert hp == 6531.0


def test_schema_to_drawio_dark_variant():
    import diagram_view
    from excel_to_mermaid import Relationship, Schema, Table

    schema = Schema(
        tables=[Table("dbo", "A"), Table("dbo", "B")],
        relationships=[Relationship("A", "B", "AID")],
    )
    light = ET.fromstring(diagram_view.schema_to_drawio(schema))
    dark = ET.fromstring(diagram_view.schema_to_drawio(schema, dark=True))

    # Dark: dark page background, dark box fill + light font, light edge stroke.
    gm_d = dark.find(".//mxGraphModel")
    assert gm_d.attrib.get("background") == "#1e1e1e"
    dv = [c.attrib["style"] for c in dark.findall(".//mxCell[@vertex='1']")]
    assert dv and all("fillColor=#2b2b2b" in s and "fontColor=#e8eaed" in s for s in dv)
    de = [c.attrib["style"] for c in dark.findall(".//mxCell[@edge='1']")]
    assert de and all("strokeColor=#9aa0a6" in s for s in de)

    # Light (default) is unchanged: white fill, no page background, no fontColor.
    gm_l = light.find(".//mxGraphModel")
    assert "background" not in gm_l.attrib
    lv = [c.attrib["style"] for c in light.findall(".//mxCell[@vertex='1']")]
    assert lv and all("fillColor=#ffffff" in s and "fontColor" not in s for s in lv)


def test_schema_to_drawio_includes_table_names_and_fk_labels():
    import diagram_view
    from excel_to_mermaid import build_schema

    records = [dict(zip(app_module.EXPECTED_HEADERS, row)) for row in ROWS]
    schema = build_schema(records)
    drawio = diagram_view.schema_to_drawio(schema)
    root = ET.fromstring(drawio)
    vertex_values = [
        cell.attrib.get("value", "") for cell in root.findall(".//mxCell[@vertex='1']")
    ]
    edge_values = [
        cell.attrib.get("value", "") for cell in root.findall(".//mxCell[@edge='1']")
    ]
    assert any("<b>Customer</b>" in value for value in vertex_values)
    assert any("<b>Order</b>" in value for value in vertex_values)
    assert "CustomerID" in edge_values


def test_schema_to_excalidraw_scene_is_valid_and_bound():
    import json

    import diagram_view
    from excel_to_mermaid import build_schema

    records = [dict(zip(app_module.EXPECTED_HEADERS, row)) for row in ROWS]
    schema = build_schema(records)
    scene = json.loads(diagram_view.schema_to_excalidraw(schema))

    assert scene["type"] == "excalidraw"
    by_id = {el["id"]: el for el in scene["elements"]}
    rects = [e for e in scene["elements"] if e["type"] == "rectangle"]
    texts = [e for e in scene["elements"] if e["type"] == "text"]
    arrows = [e for e in scene["elements"] if e["type"] == "arrow"]

    assert len(rects) == len(schema.tables)
    assert len(arrows) == len(schema.relationships)

    # Table names + FK labels present in the text elements.
    joined = "\n".join(t["text"] for t in texts)
    assert "Customer" in joined and "Order" in joined and "CustomerID" in joined

    # Every arrow is glued to real rectangles at both ends.
    for arrow in arrows:
        for side in ("startBinding", "endBinding"):
            target = arrow[side]["elementId"]
            assert target in by_id and by_id[target]["type"] == "rectangle"

    # Bound text/arrows are back-referenced from their rectangles, and all
    # referenced ids exist.
    for rect in rects:
        for ref in rect["boundElements"]:
            assert ref["id"] in by_id


def test_schema_to_excalidraw_relationship_matching_is_case_insensitive():
    import json

    import diagram_view
    from excel_to_mermaid import Relationship, Schema, Table

    schema = Schema(
        tables=[Table("dbo", "Customer"), Table("dbo", "Order")],
        relationships=[Relationship(parent_table="CUSTOMER", child_table="order", label="CustomerID")],
    )
    scene = json.loads(diagram_view.schema_to_excalidraw(schema))
    arrows = [e for e in scene["elements"] if e["type"] == "arrow"]
    assert len(arrows) == 1


def test_schema_to_excalidraw_dark_mode():
    import json

    import diagram_view
    from excel_to_mermaid import Relationship, Schema, Table

    schema = Schema(
        tables=[Table("dbo", "A"), Table("dbo", "B")],
        relationships=[Relationship(parent_table="A", child_table="B", label="X")],
    )
    # Excalidraw's dark theme inverts the canvas at render time, so a dark scene
    # keeps the normal colours and only flips the theme flag.
    scene = json.loads(diagram_view.schema_to_excalidraw(schema, dark=True))
    assert scene["appState"]["theme"] == "dark"
    assert all(e["strokeColor"] == "#1e1e1e" for e in scene["elements"])
    light = json.loads(diagram_view.schema_to_excalidraw(schema))
    assert light["appState"]["theme"] == "light"
    assert all(e["strokeColor"] == "#1e1e1e" for e in light["elements"])


def test_schema_to_excalidraw_spacing_grows_with_label_length():
    import json

    import diagram_view
    from excel_to_mermaid import Relationship, Schema, Table

    def gap(label):
        schema = Schema(
            tables=[Table("dbo", "A"), Table("dbo", "B")],
            relationships=[Relationship(parent_table="A", child_table="B", label=label)],
        )
        els = json.loads(diagram_view.schema_to_excalidraw(schema))["elements"]
        rects = {e["id"]: e for e in els if e["type"] == "rectangle"}
        a, b = rects["rect0"], rects["rect1"]
        return b["x"] - (a["x"] + a["width"])  # horizontal gap between the boxes

    # A long relationship label widens the gap so it doesn't overlap a box.
    assert gap("A_VERY_LONG_FOREIGN_KEY_COLUMN_NAME") > gap("X")


def test_schema_to_drawio_spacing_grows_with_label_length():
    import diagram_view
    from excel_to_mermaid import Relationship, Schema, Table

    def gap(label):
        schema = Schema(
            tables=[Table("dbo", "A"), Table("dbo", "B")],
            relationships=[Relationship(parent_table="A", child_table="B", label=label)],
        )
        root = ET.fromstring(diagram_view.schema_to_drawio(schema))
        geoms = {}
        for cell in root.findall(".//mxCell[@vertex='1']"):
            g = cell.find("mxGeometry")
            geoms[cell.attrib["id"]] = (float(g.attrib["x"]), float(g.attrib["width"]))
        # Two tables land side by side in the same grid row (ids "2" and "3").
        (ax, aw), (bx, _bw) = geoms["2"], geoms["3"]
        return bx - (ax + aw)  # horizontal gap between the boxes

    # A long relationship label widens the inter-box gap, same as Excalidraw.
    assert gap("A_VERY_LONG_FOREIGN_KEY_COLUMN_NAME") > gap("X")


def test_schema_to_drawio_relationship_matching_is_case_insensitive():
    import diagram_view
    from excel_to_mermaid import Relationship, Schema, Table

    schema = Schema(
        tables=[Table("dbo", "Customer"), Table("dbo", "Order")],
        relationships=[Relationship(parent_table="CUSTOMER", child_table="order", label="CustomerID")],
    )
    drawio = diagram_view.schema_to_drawio(schema)
    root = ET.fromstring(drawio)
    vertices = {
        cell.attrib["id"]: cell.attrib.get("value", "")
        for cell in root.findall(".//mxCell[@vertex='1']")
    }
    edges = root.findall(".//mxCell[@edge='1']")
    assert len(edges) == 1
    source = edges[0].attrib.get("source")
    target = edges[0].attrib.get("target")
    assert source in vertices and "<b>Customer</b>" in vertices[source]
    assert target in vertices and "<b>Order</b>" in vertices[target]


def test_schema_to_drawio_qualified_names_avoid_ambiguous_table_matches():
    import diagram_view
    from excel_to_mermaid import Relationship, Schema, Table

    schema = Schema(
        tables=[
            Table("dbo", "Customer"),
            Table("sales", "Customer"),
            Table("dbo", "Order"),
        ],
        relationships=[
            Relationship(parent_table="dbo.Customer", child_table="Order", label="CustomerID"),
            Relationship(parent_table="Customer", child_table="Order", label="AmbiguousCustomer"),
        ],
    )
    drawio = diagram_view.schema_to_drawio(schema)
    root = ET.fromstring(drawio)
    edge_values = [cell.attrib.get("value", "") for cell in root.findall(".//mxCell[@edge='1']")]
    assert "CustomerID" in edge_values
    assert "AmbiguousCustomer" not in edge_values


def test_schema_to_drawio_qualified_case_collision_is_ambiguous():
    import diagram_view
    from excel_to_mermaid import Relationship, Schema, Table

    schema = Schema(
        tables=[Table("dbo", "Customer"), Table("DBO", "customer"), Table("dbo", "Order")],
        relationships=[
            Relationship(parent_table="dbo.customer", child_table="Order", label="ShouldSkip")
        ],
    )
    drawio = diagram_view.schema_to_drawio(schema)
    root = ET.fromstring(drawio)
    edge_values = [cell.attrib.get("value", "") for cell in root.findall(".//mxCell[@edge='1']")]
    assert "ShouldSkip" not in edge_values


def _png_size(path: str) -> tuple[int, int]:
    """Return (width, height) read from a PNG header, or (0, 0) if invalid."""
    with open(path, "rb") as handle:
        header = handle.read(24)
    if len(header) < 24 or header[:8] != b"\x89PNG\r\n\x1a\n":
        return (0, 0)
    width, height = struct.unpack(">II", header[16:24])
    return (width, height)


def _ensure_sample(path: str):
    from openpyxl import Workbook

    from excel_to_mermaid import EXPECTED_HEADERS

    wb = Workbook()
    ws = wb.active
    ws.append(EXPECTED_HEADERS)
    for row in ROWS:
        ws.append(row)
    wb.save(path)


def test_window_screenshot(tmp_path):
    # Reuse a single QApplication across tests.
    app = QApplication.instance() or QApplication([])

    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))
    window.resize(1100, 760)
    assert window._options_bar.render_style().layout_direction == "LR"
    assert not window._export_btn.isHidden()
    kinds = [
        window._export_format.itemData(i)
        for i in range(window._export_format.count())
    ]
    assert kinds == ["drawio", "excalidraw", "pdf", "png", "svg"]

    out = tmp_path / "window.png"
    window.capture(str(out))

    assert out.exists(), "screenshot file was not created"
    width, height = _png_size(str(out))
    assert width > 200 and height > 200, f"screenshot looks empty: {width}x{height}"

    # Sanity: the load actually populated the model.
    assert len(window._schema.tables) == 4
    assert window._table.model().rowCount() == len(ROWS)

    del app  # keep linters quiet; app is a singleton


def test_export_drawio_checks_schema_before_prompt(monkeypatch):
    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()

    info_calls = []

    def _info(*args, **kwargs):
        info_calls.append((args, kwargs))
        return app_module.QMessageBox.Ok

    def _unexpected_dialog(*args, **kwargs):
        raise AssertionError("File dialog should not be opened when schema is empty.")

    monkeypatch.setattr(app_module.QMessageBox, "information", _info)
    monkeypatch.setattr(app_module.QFileDialog, "getSaveFileName", _unexpected_dialog)
    window._export_format.setCurrentIndex(0)  # drawio
    window.export_diagram()
    assert info_calls, "Expected an informational prompt for empty schema."

    window._diagram_view.cleanup()
    del app


def test_large_export_warns_but_does_not_block(monkeypatch):
    from excel_to_mermaid import Schema, Table

    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()

    small = Schema(tables=[Table("dbo", "A"), Table("dbo", "B")], relationships=[])
    big = Schema(
        tables=[Table("dbo", f"T{i}") for i in range(app_module.EXPORT_WARN_TABLES + 1)],
        relationships=[],
    )

    asked = []

    def _question(*args, **kwargs):
        asked.append(args)
        return app_module.QMessageBox.No

    monkeypatch.setattr(app_module.QMessageBox, "question", _question)

    # Small export: no warning, proceeds.
    assert window._confirm_large_export("excalidraw", small) is True
    assert not asked

    # Large export: warns; declining stops it (but it's a choice, not a cap).
    assert window._confirm_large_export("excalidraw", big) is False
    assert asked, "expected a confirmation prompt for a large export"

    # Accepting the warning proceeds (no hard limit).
    monkeypatch.setattr(
        app_module.QMessageBox, "question", lambda *a, **k: app_module.QMessageBox.Yes
    )
    assert window._confirm_large_export("drawio", big) is True

    window._diagram_view.cleanup()
    del app


def test_export_rendered_format_requires_webengine(monkeypatch, tmp_path):
    app = QApplication.instance() or QApplication([])
    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))
    window._diagram_view.available = False
    window._render_selection()

    info_calls = []

    def _info(*args, **kwargs):
        info_calls.append((args, kwargs))
        return app_module.QMessageBox.Ok

    def _unexpected_dialog(*args, **kwargs):
        raise AssertionError("File dialog should not open when WebEngine is unavailable.")

    monkeypatch.setattr(app_module.QMessageBox, "information", _info)
    monkeypatch.setattr(app_module.QFileDialog, "getSaveFileName", _unexpected_dialog)
    window._export_format.setCurrentIndex(2)  # pdf
    window.export_diagram()
    assert info_calls, "Expected an informational prompt when WebEngine is unavailable."

    window._diagram_view.cleanup()
    del app


def test_diagram_screenshot(tmp_path):
    """Render the actual Mermaid ER diagram to PNG and SVG (needs WebEngine)."""
    import diagram_view

    if not diagram_view.WEBENGINE_AVAILABLE:
        pytest.skip("PySide6 WebEngine not available")

    app = QApplication.instance() or QApplication([])

    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))

    png = tmp_path / "diagram.png"
    window.capture_diagram(str(png))
    assert png.exists()
    width, height = _png_size(str(png))
    assert width > 100 and height > 100, f"diagram PNG looks empty: {width}x{height}"

    svg = tmp_path / "diagram.svg"
    window.capture_diagram(str(svg))
    text = svg.read_text(encoding="utf-8")
    assert text.lstrip().startswith("<svg")
    # The rendered diagram should mention the tables from the sample.
    assert "OrderLine" in text and "Customer" in text

    # Relationship-label boxes must carry an inline style fill. Mermaid fills
    # them with an hsl() CSS rule QtSvg can't parse (it falls back to dark and
    # the label reads as a dark redaction bar in PNG/PDF); the inline style set
    # in current_svg() overrides that. See current_svg().
    import re

    boxes = re.findall(
        r'<rect[^>]*class="[^"]*relationshipLabelBox[^"]*"[^>]*>', text
    )
    assert boxes, "expected relationship-label boxes in the rendered SVG"
    assert all("style=" in b and "fill" in b for b in boxes), (
        "relationship-label boxes need an inline style fill so QtSvg doesn't "
        "render them as dark redaction bars"
    )

    # QtSvg ignores dominant-baseline:middle, so every centred text (table
    # cells, entity titles, edge labels) would ride above its intended y. It's
    # converted to an alphabetic baseline in current_svg(), so none should
    # remain in the export. See current_svg().
    assert "dominant-baseline: middle" not in text, (
        "centred text must drop dominant-baseline:middle so QtSvg positions it "
        "correctly"
    )
    assert re.search(r'class="[^"]*entityLabel[^"]*"', text), (
        "expected entity (cell) labels in the rendered SVG"
    )
    assert re.search(r'class="[^"]*relationshipLabel[^"]*"', text), (
        "expected relationship labels in the rendered SVG"
    )

    window._diagram_view.cleanup()
    del app


def _trace_path(window, start, end, *, replace):
    """Drive the From/To pickers + replace toggle, then trace the path."""
    window._selector._path_from.setCurrentText(start)
    window._selector._path_to.setCurrentText(end)
    window._selector._path_replace.setChecked(replace)
    window._find_shortest_path()


def test_trace_path_from_pickers_adds_to_selection(tmp_path):
    app = QApplication.instance() or QApplication([])

    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))
    # Customer is off the Order→OrderLine path, so it survives an additive trace.
    window._selector.clear_selection()
    window._selector.check_tables(["Customer"])

    _trace_path(window, "Order", "OrderLine", replace=False)

    selected = set(window._selector.selected_tables())
    assert {"Customer", "Order", "OrderLine"} <= selected

    window._diagram_view.cleanup()
    del app


def test_trace_path_replace_mode_discards_prior_selection(tmp_path):
    app = QApplication.instance() or QApplication([])

    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))
    window._selector.clear_selection()
    window._selector.check_tables(["Customer"])

    _trace_path(window, "Order", "OrderLine", replace=True)

    # Replace mode wipes the prior selection: only the path remains.
    assert set(window._selector.selected_tables()) == {"Order", "OrderLine"}

    window._diagram_view.cleanup()
    del app


def test_trace_path_undo_restores_prior_selection(tmp_path):
    app = QApplication.instance() or QApplication([])

    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))
    window._selector.clear_selection()
    window._selector.check_tables(["Customer"])

    _trace_path(window, "Order", "OrderLine", replace=False)
    assert window._selector._undo_btn.isEnabled()

    window._selector.undo_last_change()
    assert set(window._selector.selected_tables()) == {"Customer"}
    assert not window._selector._undo_btn.isEnabled()

    window._diagram_view.cleanup()
    del app


def test_trace_path_disconnected_keeps_selection_and_offers_no_undo(tmp_path):
    app = QApplication.instance() or QApplication([])

    sample = tmp_path / "sample.xlsx"
    _ensure_sample(str(sample))

    window = app_module.MainWindow()
    window.load_file(str(sample))
    window._selector.clear_selection()
    window._selector.check_tables(["Customer"])

    # Same table both ends is rejected before any mutation or snapshot.
    _trace_path(window, "Customer", "Customer", replace=True)
    assert set(window._selector.selected_tables()) == {"Customer"}
    assert not window._selector._undo_btn.isEnabled()

    window._diagram_view.cleanup()
    del app


def test_large_render_warning_asks_once_per_selection(monkeypatch):
    from excel_to_mermaid import Column, Schema, Table

    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    n = app_module.RENDER_WARN_LIMIT + 5  # over the warn threshold
    tables = [
        Table("dbo", f"T{i}", [Column("dbo", f"T{i}", 1, "Id", "int", is_primary_key=True)])
        for i in range(n)
    ]
    window._schema = Schema(tables=tables, relationships=[])
    window._selector.set_tables([t.name for t in tables])
    window._selector.check_all()
    # Don't actually render (keep the test fast/headless-safe).
    monkeypatch.setattr(window._diagram_view, "set_diagram", lambda *a, **k: None)

    calls = {"n": 0}

    def _q(*a, **k):
        calls["n"] += 1
        return app_module.QMessageBox.Yes

    monkeypatch.setattr(app_module.QMessageBox, "question", _q)

    window._render_selection()
    window._render_selection()  # same big selection → must not re-ask
    assert calls["n"] == 1

    # Changing the selection (still over the limit) asks again.
    window._selector.set_selected_tables([t.name for t in tables[:-1]])
    window._render_selection()
    assert calls["n"] == 2

    window._diagram_view.cleanup()
    del app


def test_empty_load_reports_no_table_data():
    from excel_to_mermaid import Schema

    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    # A file that parsed but produced no tables (no row had TableName+ColumnName).
    window._apply_loaded("weird.xlsx", [{"Foo": "bar"}], Schema(tables=[], relationships=[]), "")
    assert "no table/column data" in window._status.text().lower()

    window._diagram_view.cleanup()
    del app


def _window_with_schema(tables, rels):
    """A MainWindow with an in-memory schema loaded (no file, no render)."""
    from excel_to_mermaid import Schema, Table

    app = QApplication.instance() or QApplication([])
    window = app_module.MainWindow()
    window._schema = Schema(tables=[Table("", n) for n in tables], relationships=rels)
    window._selector.set_tables(tables)
    return app, window


def _diamond():
    """A→B, A→C, B→D, C→D — two equally short A→D paths."""
    from excel_to_mermaid import Relationship

    return ["A", "B", "C", "D"], [
        Relationship("A", "B", "ab"),
        Relationship("A", "C", "ac"),
        Relationship("B", "D", "bd"),
        Relationship("C", "D", "cd"),
    ]


def _abcd_chain():
    from excel_to_mermaid import Relationship

    return ["A", "B", "C", "D", "E"], [
        Relationship("A", "B", "ab"),
        Relationship("B", "C", "bc"),
        Relationship("C", "D", "cd"),
    ]


def _set_direction(window, value):
    combo = window._selector._path_direction
    combo.setCurrentIndex(combo.findData(value))


def test_trace_path_via_stop_forces_route(monkeypatch):
    app, window = _window_with_schema(*_diamond())
    monkeypatch.setattr(window, "_render_selection", lambda: None)

    window._selector._path_from.setCurrentText("A")
    window._selector._path_to.setCurrentText("D")
    window._selector._path_via.setCurrentText("C")  # force the A-C-D branch
    window._selector._path_replace.setChecked(True)
    window._find_shortest_path()

    assert set(window._selector.selected_tables()) == {"A", "C", "D"}

    window._diagram_view.cleanup()
    del app


def test_trace_path_popup_lets_user_choose_between_ties(monkeypatch):
    app, window = _window_with_schema(*_diamond())
    monkeypatch.setattr(window, "_render_selection", lambda: None)

    window._selector._path_from.setCurrentText("A")
    window._selector._path_to.setCurrentText("D")
    window._selector._path_replace.setChecked(True)

    seen = {}

    def _fake_get_item(_parent, _title, _label, items, _index, _editable):
        seen["items"] = list(items)
        return (items[0], True)  # sorted → "A → B → D"

    monkeypatch.setattr(app_module.QInputDialog, "getItem", _fake_get_item)
    window._find_shortest_path()

    assert seen["items"] == ["A → B → D", "A → C → D"]
    assert set(window._selector.selected_tables()) == {"A", "B", "D"}

    window._diagram_view.cleanup()
    del app


def test_trace_path_no_directed_path_offers_direction_retry(monkeypatch):
    app, window = _window_with_schema(*_abcd_chain())
    monkeypatch.setattr(window, "_render_selection", lambda: None)

    window._selector._path_from.setCurrentText("A")
    window._selector._path_to.setCurrentText("D")
    _set_direction(window, "forward")  # child→parent: A can't reach D forwards
    window._selector._path_replace.setChecked(True)

    asked = {"count": 0}

    def _yes(*_args, **_kwargs):
        asked["count"] += 1
        return app_module.QMessageBox.Yes

    monkeypatch.setattr(app_module.QMessageBox, "question", _yes)
    window._find_shortest_path()

    # Retrying without the direction constraint finds the whole chain.
    assert asked["count"] == 1
    assert set(window._selector.selected_tables()) == {"A", "B", "C", "D"}

    window._diagram_view.cleanup()
    del app


def test_trace_path_no_path_reports_and_keeps_selection(monkeypatch):
    app, window = _window_with_schema(*_abcd_chain())  # E is isolated
    monkeypatch.setattr(window, "_render_selection", lambda: None)
    window._selector.check_tables(["A"])

    window._selector._path_from.setCurrentText("A")
    window._selector._path_to.setCurrentText("E")
    window._selector._path_replace.setChecked(False)

    told = {"count": 0}
    monkeypatch.setattr(
        app_module.QMessageBox,
        "information",
        lambda *a, **k: told.__setitem__("count", told["count"] + 1),
    )
    window._find_shortest_path()

    assert told["count"] == 1
    assert set(window._selector.selected_tables()) == {"A"}  # untouched
    assert not window._selector._undo_btn.isEnabled()  # no snapshot taken

    window._diagram_view.cleanup()
    del app


def _set_related_direction(window, value):
    combo = window._selector._related_direction
    combo.setCurrentIndex(combo.findData(value))


def _hub_schema(children=12):
    """One parent H referencing `children` child tables (a fan-out hub)."""
    from excel_to_mermaid import Relationship

    kids = [f"C{i}" for i in range(children)]
    rels = [Relationship("H", kid, "fk") for kid in kids]
    return ["H", *kids], rels


def test_add_related_direction_limits_neighbours(monkeypatch):
    app, window = _window_with_schema(*_diamond())
    monkeypatch.setattr(window, "_render_selection", lambda: None)
    window._selector.clear_selection()
    window._selector.check_tables(["D"])

    # D references B and C (forward); nothing depends on D (reverse).
    _set_related_direction(window, "forward")
    window._add_related_tables()
    assert set(window._selector.selected_tables()) == {"D", "B", "C"}

    window._selector.clear_selection()
    window._selector.check_tables(["D"])
    _set_related_direction(window, "reverse")
    window._add_related_tables()
    assert set(window._selector.selected_tables()) == {"D"}  # nothing added

    window._diagram_view.cleanup()
    del app


def test_add_related_previews_and_can_be_declined(monkeypatch):
    tables, rels = _hub_schema(children=app_module.RELATED_WARN_COUNT + 2)
    app, window = _window_with_schema(tables, rels)
    monkeypatch.setattr(window, "_render_selection", lambda: None)
    window._selector.clear_selection()
    window._selector.check_tables(["H"])
    _set_related_direction(window, "reverse")  # H's dependents = all children

    asked = {"count": 0}

    def _decline(*_args, **_kwargs):
        asked["count"] += 1
        return app_module.QMessageBox.No

    monkeypatch.setattr(app_module.QMessageBox, "question", _decline)
    window._add_related_tables()

    assert asked["count"] == 1  # a big batch prompted
    assert set(window._selector.selected_tables()) == {"H"}  # declined → unchanged
    assert not window._selector._undo_btn.isEnabled()

    window._diagram_view.cleanup()
    del app


def test_add_related_accepts_and_is_undoable(monkeypatch):
    n_kids = app_module.RELATED_WARN_COUNT + 2
    tables, rels = _hub_schema(children=n_kids)
    app, window = _window_with_schema(tables, rels)
    monkeypatch.setattr(window, "_render_selection", lambda: None)
    window._selector.clear_selection()
    window._selector.check_tables(["H"])
    _set_related_direction(window, "reverse")

    monkeypatch.setattr(
        app_module.QMessageBox, "question", lambda *a, **k: app_module.QMessageBox.Yes
    )
    window._add_related_tables()

    assert len(window._selector.selected_tables()) == n_kids + 1  # H + all children
    assert window._selector._undo_btn.isEnabled()
    assert window._selector._undo_btn.text() == "&Undo add"

    window._selector.undo_last_change()
    assert set(window._selector.selected_tables()) == {"H"}
    assert not window._selector._undo_btn.isEnabled()

    window._diagram_view.cleanup()
    del app


def test_module_import_sets_webengine_flags_for_cli_mode():
    """`--screenshot-diagram` should preconfigure WebEngine flags at import time."""
    cmd = [
        sys.executable,
        "-c",
        (
            "import os, sys; "
            "sys.argv=['main.py','sample.xlsx','--screenshot-diagram','diagram.png']; "
            "import main; "
            "print(os.environ.get('QT_QPA_PLATFORM','')); "
            "print(os.environ.get('QTWEBENGINE_DISABLE_SANDBOX','')); "
            "print(os.environ.get('QTWEBENGINE_CHROMIUM_FLAGS','')); "
        ),
    ]
    env = {
        k: v
        for k, v in os.environ.items()
        if k != "QT_QPA_PLATFORM" and not k.startswith("QTWEBENGINE_")
    }
    proc = subprocess.run(
        cmd,
        cwd=os.path.dirname(__file__),
        env=env,
        capture_output=True,
        text=True,
        check=False,
    )
    assert proc.returncode == 0, proc.stderr
    lines = [line.strip() for line in proc.stdout.splitlines()]
    assert lines[:3] == [
        "offscreen",
        "1",
        "--no-sandbox --disable-gpu --in-process-gpu",
    ]


if __name__ == "__main__":
    from pathlib import Path

    with tempfile.TemporaryDirectory() as directory:
        test_window_screenshot(Path(directory))
        print("PASS  test_window_screenshot")
    with tempfile.TemporaryDirectory() as directory:
        test_diagram_screenshot(Path(directory))
        print("PASS  test_diagram_screenshot")
